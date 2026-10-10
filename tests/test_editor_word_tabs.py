"""Tab stops and baseline offsets survive the native editable round trip."""
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from pct.documento_fedele.da_docx import converti_docx
from pct.editor import html_to_docx


def test_tabbed_mixed_lines_keep_stops_styles_and_offsets(tmp_path):
    source = Document()
    p = source.add_paragraph()
    p.paragraph_format.left_indent = Pt(100)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(105))
    p.paragraph_format.tab_stops.add_tab_stop(Pt(130))
    p.add_run('\t')
    first = p.add_run('Studio')
    first.bold = True
    p.add_run('\n\t\t')
    second = p.add_run('Avvocato')
    second.italic = True
    offset = OxmlElement('w:position')
    offset.set(qn('w:val'), '7')
    second._r.get_or_add_rPr().append(offset)
    p.add_run('\n-\tIndirizzo')
    path = tmp_path / 'source.docx'
    source.save(path)
    html = ''.join(page.html for page in converti_docx(path).pagine)
    assert 'width:5pt' in html and 'width:25pt' in html
    assert '--iu-word-position:3.5' in html
    restored = Document(BytesIO(html_to_docx(html)))
    actual = restored.paragraphs[0]
    assert actual.text == p.text
    assert [t.get(qn('w:pos')) for t in actual._p.findall('w:pPr/w:tabs/w:tab', actual._p.nsmap)] == ['2100', '2600']
    assert next(r for r in actual.runs if r.text == 'Studio').bold
    r = next(r for r in actual.runs if r.text == 'Avvocato')
    assert r.italic and r._r.find(qn('w:rPr')).find(qn('w:position')).get(qn('w:val')) == '7'


def test_section_terminator_is_not_duplicated_as_an_empty_line(tmp_path):
    from docx.enum.section import WD_SECTION_START

    source = Document()
    source.add_paragraph('Prima pagina')
    source.add_section(WD_SECTION_START.NEW_PAGE)
    p = source.add_paragraph('Seconda pagina')
    p.paragraph_format.space_before = Pt(36)
    path = tmp_path / 'sections.docx'
    source.save(path)
    html = ''.join(page.html for page in converti_docx(path).pagine)
    restored = Document(BytesIO(html_to_docx(html)))
    assert len(restored.sections) == 2
    assert [p.text for p in restored.paragraphs] == [p.text for p in source.paragraphs]
    assert restored.paragraphs[-1].paragraph_format.space_before.pt == 36


def test_browser_serialized_margin_shorthand_preserves_native_geometry():
    html = '<p style="margin:34.35pt 187.2pt 0pt 196.9pt;text-indent:0pt;line-height:14.1pt">Timbro</p>'
    p = Document(BytesIO(html_to_docx(html))).paragraphs[0]
    assert p.paragraph_format.space_before.pt == 34.35
    assert p.paragraph_format.space_after.pt == 0
    assert p.paragraph_format.left_indent.pt == 196.9
    assert p.paragraph_format.right_indent.pt == 187.2
    assert p.paragraph_format.line_spacing.pt == 14.1
    html = '<p style="margin-left:100pt;margin:12px 20pt;margin-left:30pt">Corpo</p>'
    p = Document(BytesIO(html_to_docx(html))).paragraphs[0]
    assert p.paragraph_format.left_indent.pt == 30
    assert p.paragraph_format.right_indent.pt == 20
    assert p.paragraph_format.space_before.pt == 9


def test_picked_text_color_and_highlight_persist_with_mixed_styles(tmp_path):
    html = '<p><strong>Prima <span style="color:rgb(0,85,170);background-color:rgb(254,243,199)"><em>CHIEDE</em></span> dopo</strong></p>'
    raw = html_to_docx(html)
    doc = Document(BytesIO(raw))
    chosen = next(r for r in doc.paragraphs[0].runs if r.text == 'CHIEDE')
    assert str(chosen.font.color.rgb) == '0055AA'
    assert chosen.bold and chosen.italic
    assert chosen._r.find(qn('w:rPr')).find(qn('w:shd')).get(qn('w:fill')) == 'fef3c7'
    assert all(r.font.color.rgb is None for r in doc.paragraphs[0].runs if r.text != 'CHIEDE')
    path = tmp_path / 'color.docx'
    path.write_bytes(raw)
    imported = ''.join(p.html for p in converti_docx(path).pagine)
    restored = Document(BytesIO(html_to_docx(imported)))
    chosen = next(r for r in restored.paragraphs[0].runs if r.text == 'CHIEDE')
    assert str(chosen.font.color.rgb) == '0055AA'
    assert chosen.bold and chosen.italic
    assert chosen._r.find(qn('w:rPr')).find(qn('w:shd')).get(qn('w:fill')) == 'fef3c7'
