"""Il DOCX letto direttamente, senza passare per un convertitore esterno.

Un atto che torna dall'editor, viene modificato in Word e rientra deve
ritrovare quello che aveva: carattere, corpo, colore, sottolineature,
allineamenti, rientri, interlinea, elenchi, tabelle con intestazioni e sfondi,
formato e margini della pagina.

Prima questa strada aveva due varianti, tutte e due con dei buchi. La rotta di
importazione usava `mammoth`, che tiene grassetto, corsivo, tabelle ed elenchi
e butta via tutto il resto — misurato su un atto: sottolineato, colore,
allineamenti, rientri, carattere, corpo e margini. `converti_docx` chiamava
LibreOffice per fare DOCX -> PDF e poi rileggeva il PDF: fedele, ma richiede un
programma da mezzo giga sul server e un processo esterno con il suo timeout, e
sul server di IUSENTRA quel programma non c'e'.

Qui il DOCX si legge com'e', con python-docx, e si scrive nell'HTML
dell'editor riusando `_html_tratti`: la stessa resa dei documenti che arrivano
da un PDF, cosi' un atto ha lo stesso aspetto da qualunque parte entri.

Solo python-docx (MIT). Nessun programma esterno, nessuna libreria AGPL.
"""

from __future__ import annotations

import base64
from pathlib import Path

from docx import Document as ApriDocx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor

from .modello import DocumentoConvertito, PaginaConvertita, Tratto
from .paragrafi import _html_tratti
from .taratura import _pt, pila_font

#: Gli allineamenti di Word, tradotti in quelli del foglio di stile.
ALLINEAMENTI = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}

#: Corpo e carattere da usare quando il documento non li dichiara da nessuna
#: parte: sono quelli che Word stesso mette di suo.
CORPO_PREDEFINITO = 11.0
CARATTERE_PREDEFINITO = "Calibri"

#: Un paragrafo con questo stile e' una voce di elenco.
STILI_ELENCO_NUMERATO = ("list number", "elenco numerato")
STILI_ELENCO_PUNTATO = ("list bullet", "list paragraph", "elenco puntato")

NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class DocxError(ValueError):
    """Il documento non si apre o non e' un DOCX."""


# ---------------------------------------------------------------------------
# Quello che Word non scrive sul pezzo ma eredita dallo stile
# ---------------------------------------------------------------------------

def _eredita(oggetto, attributo: str, paragrafo, documento):
    """Il valore dichiarato sul tratto, o quello che eredita.

    In un DOCX quasi niente e' scritto dove lo si cerca: il corpo di una parola
    puo' stare sul tratto, sullo stile del paragrafo, sullo stile da cui quello
    discende, o solo nelle impostazioni del documento. Chi legge solo il primo
    livello trova `None` quasi sempre — ed e' il motivo per cui le conversioni
    veloci perdono carattere e corpo.
    """
    valore = getattr(oggetto, attributo, None)
    if valore is not None:
        return valore

    stile = getattr(paragrafo, "style", None)
    visti = 0
    while stile is not None and visti < 8:
        carattere = getattr(stile, "font", None)
        valore = getattr(carattere, attributo, None) if carattere is not None else None
        if valore is not None:
            return valore
        stile = getattr(stile, "base_style", None)
        visti += 1

    try:
        normale = documento.styles["Normal"].font
        return getattr(normale, attributo, None)
    except (KeyError, AttributeError):
        return None


def _formato_ereditato(paragrafo, attributo):
    """Il valore del formato scritto sul paragrafo, o quello che eredita.

    Un atto scritto con gli stili di Word — cioe' quasi ogni atto — non dichiara
    niente sul paragrafo: il giustificato, il rientro e la spaziatura stanno
    sullo stile. Chi legge solo `paragraph.paragraph_format` trova `None` e il
    documento torna tutto allineato a sinistra.
    """
    valore = getattr(paragrafo.paragraph_format, attributo, None)
    if valore is not None:
        return valore
    stile = getattr(paragrafo, "style", None)
    visti = 0
    while stile is not None and visti < 8:
        formato = getattr(stile, "paragraph_format", None)
        valore = getattr(formato, attributo, None) if formato is not None else None
        if valore is not None:
            return valore
        stile = getattr(stile, "base_style", None)
        visti += 1
    return None


def _colore(tratto, paragrafo, documento) -> str:
    valore = getattr(getattr(tratto, "font", None), "color", None)
    rgb = getattr(valore, "rgb", None) if valore is not None else None
    if rgb is None:
        stile = getattr(paragrafo, "style", None)
        visti = 0
        while stile is not None and visti < 8 and rgb is None:
            carattere = getattr(stile, "font", None)
            colore_stile = getattr(carattere, "color", None) if carattere is not None else None
            rgb = getattr(colore_stile, "rgb", None) if colore_stile is not None else None
            stile = getattr(stile, "base_style", None)
            visti += 1
    if rgb is None or not isinstance(rgb, RGBColor):
        return "#000000"
    return f"#{rgb}".lower()


def _evidenziazione(tratto) -> str | None:
    """Il colore dell'evidenziatore, quando c'e'."""
    nome = getattr(getattr(tratto, "font", None), "highlight_color", None)
    if nome is None:
        return None
    tinte = {
        "YELLOW": "#ffff00", "BRIGHT_GREEN": "#00ff00", "TURQUOISE": "#00ffff",
        "PINK": "#ff00ff", "RED": "#ff0000", "GRAY_25": "#c0c0c0",
        "GRAY_50": "#808080", "DARK_YELLOW": "#808000", "TEAL": "#008080",
        "BLUE": "#0000ff", "GREEN": "#008000", "VIOLET": "#800080",
    }
    return tinte.get(str(getattr(nome, "name", nome)).upper())


def _sfondo_cella(cella) -> str | None:
    """Il colore di riempimento di una cella, letto dal suo XML."""
    try:
        ombra = cella._tc.tcPr.find(f"{NS}shd")
    except AttributeError:
        return None
    if ombra is None:
        return None
    riempimento = ombra.get(f"{NS}fill")
    if not riempimento or riempimento in ("auto", "FFFFFF", "ffffff"):
        return None
    return f"#{riempimento}".lower()


def _unione(cella) -> tuple[int, bool]:
    """(quante colonne occupa, se prosegue una cella di sopra)."""
    colonne, prosegue = 1, False
    try:
        proprieta = cella._tc.tcPr
    except AttributeError:
        return colonne, prosegue
    if proprieta is None:
        return colonne, prosegue
    esteso = proprieta.find(f"{NS}gridSpan")
    if esteso is not None:
        try:
            colonne = max(1, int(esteso.get(f"{NS}val")))
        except (TypeError, ValueError):
            pass
    verticale = proprieta.find(f"{NS}vMerge")
    if verticale is not None and (verticale.get(f"{NS}val") or "continue") == "continue":
        prosegue = True
    return colonne, prosegue


# ---------------------------------------------------------------------------
# Dal paragrafo del DOCX ai tratti del modello
# ---------------------------------------------------------------------------

def _pezzi_del_paragrafo(paragrafo):
    """I tratti del paragrafo, con il loro collegamento quando ce l'hanno.

    `paragraph.runs` salta quello che sta dentro un `w:hyperlink`: su un atto
    che cita una PEC o un riferimento normativo, l'indirizzo non perdeva solo
    il collegamento — spariva anche il testo.
    """
    from docx.text.run import Run

    for elemento in paragrafo._p.iterchildren():
        if elemento.tag == f"{NS}r":
            yield Run(elemento, paragrafo), None
        elif elemento.tag == f"{NS}hyperlink":
            indirizzo = _indirizzo_collegamento(elemento, paragrafo)
            for interno in elemento.iterchildren(f"{NS}r"):
                yield Run(interno, paragrafo), indirizzo


def _indirizzo_collegamento(elemento, paragrafo) -> str | None:
    """L'indirizzo di un `w:hyperlink`, risolto nelle relazioni del documento."""
    RIF = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    riferimento = elemento.get(RIF)
    if not riferimento:
        ancora = elemento.get(f"{NS}anchor")
        return f"#{ancora}" if ancora else None
    try:
        return paragrafo.part.rels[riferimento].target_ref
    except (KeyError, AttributeError):
        return None


def _interruzioni(pezzo) -> tuple[bool, bool]:
    """(c'e' un a capo, c'e' un salto di pagina) dentro questo tratto."""
    a_capo = salto = False
    for interruzione in pezzo._r.iterchildren(f"{NS}br"):
        if interruzione.get(f"{NS}type") == "page":
            salto = True
        else:
            a_capo = True
    return a_capo, salto


def _tratti(paragrafo, documento, corpo_base: float,
            nomi: set[str] | None = None) -> list[Tratto]:
    fuori: list[Tratto] = []
    for pezzo, indirizzo in _pezzi_del_paragrafo(paragrafo):
        testo = pezzo.text
        a_capo, _salto = _interruzioni(pezzo)
        if a_capo:
            testo = f"{testo}\n"
        if not testo:
            continue
        misura = _eredita(pezzo.font, "size", paragrafo, documento)
        corpo = round(misura.pt, 1) if misura is not None else corpo_base
        nome = _eredita(pezzo.font, "name", paragrafo, documento) or CARATTERE_PREDEFINITO
        if nomi is not None and testo.strip():
            nomi.add(str(nome))
        fuori.append(Tratto(
            testo=testo,
            famiglia=pila_font(str(nome)),
            corpo=corpo,
            grassetto=bool(_eredita(pezzo, "bold", paragrafo, documento)),
            corsivo=bool(_eredita(pezzo, "italic", paragrafo, documento)),
            sottolineato=bool(_eredita(pezzo, "underline", paragrafo, documento)),
            barrato=bool(getattr(pezzo.font, "strike", None)),
            apice=bool(getattr(pezzo.font, "superscript", None)),
            pedice=bool(getattr(pezzo.font, "subscript", None)),
            colore=_colore(pezzo, paragrafo, documento),
            evidenziato=_evidenziazione(pezzo),
            collegamento=indirizzo,
        ))
    return fuori


def _ha_salto_di_pagina(paragrafo) -> bool:
    """Vero se il paragrafo porta un salto di pagina."""
    for pezzo in paragrafo._p.iterchildren(f"{NS}r"):
        for interruzione in pezzo.iterchildren(f"{NS}br"):
            if interruzione.get(f"{NS}type") == "page":
                return True
    return False


def _con_a_capo(html: str) -> str:
    """Gli a capo dentro un paragrafo diventano `<br>`.

    Nel testo arrivano come ritorni a riga, e l'HTML li collasserebbe in uno
    spazio: due righe di un indirizzo diventerebbero una sola.
    """
    return html.replace("\n", "<br>")


def _stile_paragrafo(paragrafo) -> list[str]:
    """Allineamento, rientri, interlinea e spazi, come li ha lasciati Word."""
    stile: list[str] = []

    allineamento = ALLINEAMENTI.get(_formato_ereditato(paragrafo, "alignment"))
    if allineamento and allineamento != "left":
        stile.append(f"text-align:{allineamento}")

    for attributo, proprieta in (("left_indent", "margin-left"),
                                 ("right_indent", "margin-right"),
                                 ("first_line_indent", "text-indent")):
        misura = _formato_ereditato(paragrafo, attributo)
        if misura is not None and abs(misura.pt) > 0.5:
            stile.append(f"{proprieta}:{_pt(misura.pt)}pt")

    for attributo, proprieta in (("space_before", "margin-top"),
                                 ("space_after", "margin-bottom")):
        misura = _formato_ereditato(paragrafo, attributo)
        if misura is not None and misura.pt > 0.5:
            stile.append(f"{proprieta}:{_pt(misura.pt)}pt")

    interlinea = _formato_ereditato(paragrafo, "line_spacing")
    if interlinea is not None:
        if isinstance(interlinea, float):
            if abs(interlinea - 1.0) > 0.05:
                stile.append(f"line-height:{interlinea:.2f}")
        elif getattr(interlinea, "pt", None):
            stile.append(f"line-height:{_pt(interlinea.pt)}pt")

    return stile


def _livello_elenco(paragrafo) -> int:
    """Il livello di annidamento della voce: 0 e' il primo."""
    try:
        numerazione = paragrafo._p.pPr.numPr
    except AttributeError:
        return 0
    if numerazione is None:
        return 0
    livello = numerazione.find(f"{NS}ilvl")
    if livello is None:
        return 0
    try:
        return max(0, min(8, int(livello.get(f"{NS}val"))))
    except (TypeError, ValueError):
        return 0


def _voce_di_elenco(paragrafo) -> str | None:
    """`ol`, `ul`, o niente se il paragrafo non e' una voce di elenco."""
    nome = (getattr(getattr(paragrafo, "style", None), "name", "") or "").lower()
    if any(n in nome for n in STILI_ELENCO_NUMERATO):
        return "ol"
    if any(n in nome for n in STILI_ELENCO_PUNTATO):
        return "ul"
    return None


def _tutti_i_paragrafi(documento):
    """Ogni paragrafo del documento, comprese le celle delle tabelle."""
    yield from documento.paragraphs
    for tabella in documento.tables:
        for riga in tabella.rows:
            for cella in riga.cells:
                yield from cella.paragraphs


def _corpo_prevalente(documento) -> tuple[float, str]:
    """Corpo e carattere usati piu' spesso: fanno da riferimento alla pagina.

    Si pesa per quante lettere sono scritte con quella combinazione, non per
    quanti paragrafi la usano: un titolo e' un paragrafo come il corpo
    dell'atto, ma non e' lui a dare il tono al documento.
    """
    conteggio: dict[tuple[float, str], int] = {}
    for paragrafo in _tutti_i_paragrafi(documento):
        for pezzo in paragrafo.runs:
            if not pezzo.text.strip():
                continue
            misura = _eredita(pezzo.font, "size", paragrafo, documento)
            nome = _eredita(pezzo.font, "name", paragrafo, documento) or CARATTERE_PREDEFINITO
            chiave = (round(misura.pt, 1) if misura is not None else CORPO_PREDEFINITO, str(nome))
            conteggio[chiave] = conteggio.get(chiave, 0) + len(pezzo.text)
    if not conteggio:
        return CORPO_PREDEFINITO, pila_font(CARATTERE_PREDEFINITO)
    corpo, nome = max(conteggio.items(), key=lambda voce: voce[1])[0]
    return corpo, pila_font(nome)


# ---------------------------------------------------------------------------
# Tabelle
# ---------------------------------------------------------------------------

def _immagini(paragrafo, documento) -> list[str]:
    """Le immagini del paragrafo, come `<img>` con i dati dentro.

    Logo dello studio, firma scansionata, timbro: senza queste un atto
    importato perde l'intestazione, e l'avvocato se ne accorge solo davanti
    alla stampa.
    """
    RIF = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    from .immagini import alleggerisci_immagine

    fuori: list[str] = []
    for disegno in paragrafo._p.iter():
        if not disegno.tag.endswith("}blip"):
            continue
        riferimento = disegno.get(RIF)
        if not riferimento:
            continue
        try:
            parte = documento.part.related_parts[riferimento]
            dati = parte.blob
        except (KeyError, AttributeError):
            continue
        if not dati:
            continue
        estensione = (getattr(parte, "partname", "").rsplit(".", 1)[-1] or "png").lower()
        try:
            dati, estensione = alleggerisci_immagine(dati, estensione)
        except Exception:
            estensione = estensione or "png"
        tipo = "jpeg" if estensione in ("jpg", "jpeg") else estensione
        sorgente = f"data:image/{tipo};base64," + base64.b64encode(dati).decode()
        larghezza = _larghezza_immagine(disegno)
        misura = f";width:{_pt(larghezza)}pt" if larghezza else ""
        fuori.append(
            f'<img src="{sorgente}" style="max-width:100%;height:auto{misura}" '
            f'alt="Immagine del documento">'
        )
    return fuori


def _larghezza_immagine(blip) -> float | None:
    """La larghezza dichiarata nel disegno, in punti."""
    for antenato in blip.iterancestors():
        for estensione in antenato.iter():
            if estensione.tag.endswith("}extent"):
                try:
                    return int(estensione.get("cx")) / 12700.0   # EMU -> punti
                except (TypeError, ValueError):
                    return None
    return None


def _e_intestazione(tabella) -> bool:
    """Vero se la prima riga della tabella e' un'intestazione.

    Word lo dichiara con `tblHeader` sulle righe che si ripetono a ogni pagina.
    Quando non lo dichiara vale la regola degli atti: la prima riga ha uno
    sfondo suo, oppure e' tutta in grassetto e le altre no.
    """
    if not tabella.rows:
        return False
    prima = tabella.rows[0]
    try:
        proprieta = prima._tr.trPr
        if proprieta is not None and proprieta.find(f"{NS}tblHeader") is not None:
            return True
    except AttributeError:
        pass

    if any(_sfondo_cella(c) for c in prima.cells):
        return True

    def _tutta_grassetto(riga) -> bool:
        pezzi = [p for c in riga.cells for par in c.paragraphs
                 for p in par.runs if p.text.strip()]
        return bool(pezzi) and all(p.bold for p in pezzi)

    if not _tutta_grassetto(prima):
        return False
    return not any(_tutta_grassetto(r) for r in tabella.rows[1:])


def _larghezze_colonne(tabella) -> list[float]:
    """La larghezza di ogni colonna in percentuale, come la dichiara il DOCX.

    Senza questa, l'editor spartisce lo spazio in parti uguali: un prospetto
    con la descrizione larga e l'importo stretto esce sbilenco.
    """
    try:
        griglia = tabella._tbl.find(f"{NS}tblGrid")
    except AttributeError:
        return []
    if griglia is None:
        return []
    misure = []
    for colonna in griglia.iterchildren(f"{NS}gridCol"):
        try:
            misure.append(float(colonna.get(f"{NS}w")))
        except (TypeError, ValueError):
            misure.append(0.0)
    totale = sum(misure)
    if totale <= 0:
        return []
    return [m / totale * 100 for m in misure]


def _ha_bordi(tabella) -> bool:
    """Vero se la tabella dichiara dei bordi visibili."""
    try:
        proprieta = tabella._tbl.tblPr
    except AttributeError:
        return False
    if proprieta is None:
        return False
    bordi = proprieta.find(f"{NS}tblBorders")
    if bordi is None:
        nome = (getattr(getattr(tabella, "style", None), "name", "") or "").lower()
        return "grid" in nome or "griglia" in nome
    for lato in bordi.iterchildren():
        valore = lato.get(f"{NS}val")
        if valore and valore not in ("none", "nil"):
            return True
    return False


def _html_tabella(tabella, documento, corpo_base: float, famiglia_base: str,
                  nomi: set[str] | None = None) -> str:
    intestazione = _e_intestazione(tabella)
    larghezze = _larghezze_colonne(tabella)
    classi = "iu-doc-tabella"
    if not _ha_bordi(tabella):
        classi += " iu-doc-tabella--senza-bordi"
    pezzi = [f'<table class="{classi}"><tbody>']
    for indice_riga, riga in enumerate(tabella.rows):
        pezzi.append("<tr>")
        vista: set = set()
        indice_colonna = 0
        for cella in riga.cells:
            if id(cella._tc) in vista:      # una cella unita compare piu' volte
                continue
            vista.add(id(cella._tc))
            colonne, prosegue = _unione(cella)
            inizio_colonna = indice_colonna
            indice_colonna += colonne
            if prosegue:
                continue

            marchio = "th" if (intestazione and indice_riga == 0) else "td"
            attributi = f' colspan="{colonne}"' if colonne > 1 else ""

            stile = []
            if larghezze:
                quota = sum(larghezze[inizio_colonna:inizio_colonna + colonne])
                if quota > 0:
                    stile.append(f"width:{quota:.1f}%")
            sfondo = _sfondo_cella(cella)
            if sfondo:
                stile.append(f"background-color:{sfondo}")

            contenuto = []
            for paragrafo in cella.paragraphs:
                tratti = _tratti(paragrafo, documento, corpo_base, nomi)
                if not tratti:
                    continue
                allineamento = ALLINEAMENTI.get(_formato_ereditato(paragrafo, "alignment"))
                interno = _con_a_capo(_html_tratti(tratti, corpo_base, famiglia_base))
                if allineamento and allineamento != "left":
                    interno = f'<span style="display:block;text-align:{allineamento}">{interno}</span>'
                contenuto.append(interno)

            attributo_stile = f' style="{";".join(stile)}"' if stile else ""
            pezzi.append(
                f"<{marchio}{attributi}{attributo_stile}>"
                f'{"<br>".join(contenuto) or "&nbsp;"}</{marchio}>'
            )
        pezzi.append("</tr>")
    pezzi.append("</tbody></table>")
    return "".join(pezzi)


# ---------------------------------------------------------------------------
# Il documento intero
# ---------------------------------------------------------------------------

def _in_ordine(documento):
    """Paragrafi e tabelle nell'ordine in cui stanno nel documento.

    `document.paragraphs` e `document.tables` sono due elenchi separati: chi
    usa quelli si ritrova tutte le tabelle in fondo, staccate dal testo che
    le introduce.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for elemento in documento.element.body.iterchildren():
        if elemento.tag == f"{NS}p":
            yield Paragraph(elemento, documento)
        elif elemento.tag == f"{NS}tbl":
            yield Table(elemento, documento)


def _formato_pagina(documento) -> dict:
    if not documento.sections:
        return {}
    sezione = documento.sections[0]

    def _punti(misura, altrimenti: float) -> float:
        return round(misura.pt, 1) if misura is not None else altrimenti

    larghezza = _punti(sezione.page_width, 595.3)
    altezza = _punti(sezione.page_height, 841.9)
    corto, lungo = min(larghezza, altezza), max(larghezza, altezza)
    nome = f"{corto:.0f}×{lungo:.0f} pt"
    for etichetta, (fl, fa) in {
        "A4": (595.3, 841.9), "A5": (419.5, 595.3), "A3": (841.9, 1190.6),
        "Letter": (612.0, 792.0), "Legal": (612.0, 1008.0),
    }.items():
        if abs(corto - fl) < 6 and abs(lungo - fa) < 6:
            nome = etichetta
            break

    return {
        "formato": nome,
        "orientamento": "orizzontale" if larghezza > altezza else "verticale",
        "larghezza_pt": larghezza,
        "altezza_pt": altezza,
        "margini_pt": {
            "alto": _punti(sezione.top_margin, 56.7),
            "destro": _punti(sezione.right_margin, 56.7),
            "basso": _punti(sezione.bottom_margin, 56.7),
            "sinistro": _punti(sezione.left_margin, 56.7),
        },
    }


def converti_docx(percorso: str | Path) -> DocumentoConvertito:
    """Legge un DOCX e lo consegna all'editor come fa `converti` con un PDF."""
    try:
        documento = ApriDocx(str(percorso))
    except Exception as errore:
        raise DocxError(f"documento illeggibile: {Path(percorso).name}") from errore

    corpo_base, famiglia_base = _corpo_prevalente(documento)
    formato = _formato_pagina(documento)

    pagine: list[list[str]] = [[]]
    aperti: list[str] = []          # gli elenchi aperti, uno per livello
    caratteri: set[str] = set()
    elementi = 0
    tabelle = 0
    immagini = 0

    def pezzi() -> list[str]:
        return pagine[-1]

    def _chiudi_elenchi(fino_a: int = 0):
        while len(aperti) > fino_a:
            pezzi().append(f"</{aperti.pop()}>")

    for blocco in _in_ordine(documento):
        if blocco.__class__.__name__ == "Table":
            _chiudi_elenchi()
            pezzi().append(_html_tabella(blocco, documento, corpo_base, famiglia_base, caratteri))
            tabelle += 1
            continue

        if _ha_salto_di_pagina(blocco):
            _chiudi_elenchi()
            pagine.append([])

        figure = _immagini(blocco, documento)
        tratti = _tratti(blocco, documento, corpo_base, caratteri)
        scritto = "".join(t.testo for t in tratti).strip()
        if not scritto and not figure:
            _chiudi_elenchi()
            continue

        if figure:
            _chiudi_elenchi()
            allineamento = ALLINEAMENTI.get(_formato_ereditato(blocco, "alignment")) or "left"
            pezzi().append(
                f'<p style="text-align:{allineamento};margin:0.35em 0">{"".join(figure)}</p>'
            )
            immagini += len(figure)
            elementi += 1
            if not scritto:
                continue

        interno = _con_a_capo(_html_tratti(tratti, corpo_base, famiglia_base))
        marchio = _voce_di_elenco(blocco)
        if marchio:
            livello = _livello_elenco(blocco) + 1
            _chiudi_elenchi(livello)
            while len(aperti) < livello:
                pezzi().append(f"<{marchio}>")
                aperti.append(marchio)
            if aperti[-1] != marchio:
                _chiudi_elenchi(livello - 1)
                pezzi().append(f"<{marchio}>")
                aperti.append(marchio)
            pezzi().append(f"<li>{interno}</li>")
        else:
            _chiudi_elenchi()
            stile = _stile_paragrafo(blocco)
            attributo = f' style="{";".join(stile)}"' if stile else ""
            pezzi().append(f"<p{attributo}>{interno}</p>")
        elementi += 1

    _chiudi_elenchi()

    stile_pagina = [f"font-family:{famiglia_base}", f"font-size:{_pt(corpo_base)}pt"]
    html = "".join(
        f'<section class="iu-doc-pagina" data-pagina="{numero}" data-origine="docx" '
        f'style="{";".join(stile_pagina)}">{"".join(contenuto)}</section>'
        for numero, contenuto in enumerate(pagine, start=1)
    )

    esito = DocumentoConvertito()
    esito.html = html
    esito.formato = formato
    esito.caratteri = sorted(c for c in caratteri if c)
    esito.pagine = [PaginaConvertita(
        numero=numero,
        html=(f'<section class="iu-doc-pagina" data-pagina="{numero}" data-origine="docx" '
              f'style="{";".join(stile_pagina)}">{"".join(contenuto)}</section>'),
        larghezza=formato.get("larghezza_pt", 595.3),
        altezza=formato.get("altezza_pt", 841.9),
        orientamento=formato.get("orientamento", "verticale"),
        margini=(
            formato.get("margini_pt", {}).get("alto", 56.7),
            formato.get("margini_pt", {}).get("destro", 56.7),
            formato.get("margini_pt", {}).get("basso", 56.7),
            formato.get("margini_pt", {}).get("sinistro", 56.7),
        ),
        elementi=elementi,
        tabelle=tabelle,
        immagini=immagini,
    ) for numero, contenuto in enumerate(pagine, start=1)]
    if not elementi and not tabelle:
        esito.avvisi.append("il documento non contiene testo")
    return esito
