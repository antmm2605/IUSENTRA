"""Conversione del documento dell'editor in `.docx`: la forma deve sopravvivere.

Quando l'avvocato salva o scarica un documento dall'editor, il `.docx` deve
essere il documento che ha davanti: titoli che restano titoli, grassetto e
corsivo dove li ha messi, elenchi che restano elenchi, tabelle che restano
tabelle, formule centrate e sottoscrizioni a destra.

Difetto trovato il 14/09/2026 e qui bloccato per sempre: l'HTML veniva avvolto
in un `<div>` e il convertitore trattava quel contenitore come un capoverso,
quindi **ogni** documento salvato dall'editor usciva come un unico paragrafo
senza titoli, senza grassetto, senza elenchi e senza tabelle.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from pct.editor import html_to_docx

ATTO = (
    "<h1>ATTO DI CITAZIONE</h1>"
    "<p>Il sottoscritto avvocato espone quanto segue.</p>"
    "<p><strong>Grassetto</strong>, <em>corsivo</em> e <strong><em>entrambi</em></strong>.</p>"
    '<p style="text-align:center"><strong>P.Q.M.</strong></p>'
    "<ul><li>Prima voce</li><li>Seconda voce</li></ul>"
    "<ol><li>Voce numerata</li></ol>"
    "<table><tr><th>Voce</th><th>Importo</th></tr><tr><td>Diritti</td><td>100,00</td></tr></table>"
    '<p style="text-align:right">Avv. Mario Rossi</p>'
)


@pytest.fixture(scope="module")
def documento():
    return Document(io.BytesIO(html_to_docx(ATTO, "Atto di citazione")))


def _paragrafi(documento):
    return [par for par in documento.paragraphs if par.text.strip()]


def test_il_documento_non_si_appiattisce_in_un_unico_paragrafo(documento):
    """Il difetto originale: tutto il documento in un paragrafo solo."""
    assert len(_paragrafi(documento)) >= 7


def test_i_titoli_restano_titoli(documento):
    primo = _paragrafi(documento)[0]
    assert primo.style.name == "Heading 1"
    assert primo.text == "ATTO DI CITAZIONE"


def test_grassetto_e_corsivo_sopravvivono_anche_annidati(documento):
    riga = next(par for par in _paragrafi(documento) if "Grassetto" in par.text)
    testi = {run.text.strip(): (bool(run.bold), bool(run.italic)) for run in riga.runs if run.text.strip()}
    assert testi.get("Grassetto") == (True, False)
    assert testi.get("corsivo") == (False, True)
    assert testi.get("entrambi") == (True, True), "il corsivo dentro il grassetto e' andato perso"


def test_l_allineamento_del_capoverso_arriva_nel_documento(documento):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    centrato = next(par for par in _paragrafi(documento) if par.text.strip() == "P.Q.M.")
    destra = next(par for par in _paragrafi(documento) if "Rossi" in par.text)
    assert centrato.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert destra.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_gli_elenchi_restano_elenchi(documento):
    stili = [par.style.name for par in _paragrafi(documento)]
    assert "List Bullet" in stili
    assert "List Number" in stili
    puntati = [par.text for par in _paragrafi(documento) if par.style.name == "List Bullet"]
    assert puntati == ["Prima voce", "Seconda voce"]


def test_le_tabelle_restano_tabelle(documento):
    assert len(documento.tables) == 1
    tabella = documento.tables[0]
    assert [cella.text for cella in tabella.rows[0].cells] == ["Voce", "Importo"]
    assert [cella.text for cella in tabella.rows[1].cells] == ["Diritti", "100,00"]


def test_il_testo_non_si_incolla_fra_un_blocco_e_il_successivo(documento):
    """Sintomo del difetto: «ATTO DI CITAZIONEIl sottoscritto…» in un blocco solo."""
    for paragrafo in _paragrafi(documento):
        assert "CITAZIONEIl" not in paragrafo.text


def test_un_documento_senza_contenitore_si_converte_ugualmente():
    documento = Document(io.BytesIO(html_to_docx("<h2>Memoria</h2><p>Testo.</p>", "Memoria")))
    stili = [(par.style.name, par.text) for par in documento.paragraphs if par.text.strip()]
    assert stili == [("Heading 2", "Memoria"), ("Normal", "Testo.")]


def test_un_html_non_valido_non_fa_perdere_il_testo():
    documento = Document(io.BytesIO(html_to_docx("<p>Testo senza chiusura", "Prova")))
    assert any("Testo senza chiusura" in par.text for par in documento.paragraphs)


def test_l_interruzione_di_pagina_dell_editor_diventa_interruzione_nel_documento():
    """Le pagine dell'originale restano pagine: stesso marcatore di editor e PDF."""
    html = '<p>Prima pagina</p><hr class="iu-ted-page-break" data-iu-page-break="true"><p>Seconda pagina</p>'
    documento = Document(io.BytesIO(html_to_docx(html, "Prova")))
    assert _interruzioni_di_pagina(documento) == 1


def test_una_linea_orizzontale_qualunque_non_spezza_la_pagina():
    documento = Document(io.BytesIO(html_to_docx("<p>Testo</p><hr><p>Altro testo</p>", "Prova")))
    assert _interruzioni_di_pagina(documento) == 0
    assert [par.text for par in _paragrafi(documento)] == ["Testo", "Altro testo"]


def _interruzioni_di_pagina(documento) -> int:
    marca = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return sum(
        1
        for paragrafo in documento.paragraphs
        for run in paragrafo.runs
        for interruzione in run._element.findall(f".//{marca}br")
        if interruzione.get(f"{marca}type") == "page"
    )
