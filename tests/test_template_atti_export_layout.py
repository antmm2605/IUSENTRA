"""Export RTF, DOCX e PDF fedeli all'impaginazione dell'editor atti."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT

from pct.template_atti_page_setup import page_setup_from_layout
from pct.template_atti_rtf import build_rtf_document, rtf_escape
from web.services.template_atti_export import build_docx_export, build_rtf_export


class _Timbro:
    enabled = True

    def to_lines(self):
        return [{"text": "STUDIO LEGALE ROSSI", "bold": True}, {"text": "Via Roma 1 - 00100 Roma"}]


EDITOR_HTML = (
    "<h1>Atto di citazione</h1>"
    "<p style=\"text-align: center\">Città <strong>grassetto</strong> <s>barrato</s><br>a capo</p>"
    "<ul><li>primo</li><li>secondo</li></ul>"
    "<hr class=\"iu-ted-page-break\" data-iu-page-break=\"true\">"
    "<p>Pagina successiva</p>"
)


def test_page_setup_riproduce_margini_e_testo_sotto_il_timbro():
    setup = page_setup_from_layout({"page_orientation": "orizzontale", "margin_top_mm": 20, "stamp_offset_y_mm": 4, "stamp_anchor": "page_margin"})

    assert setup.landscape is True
    assert (setup.width_mm, setup.height_mm) == (297.0, 210.0)
    assert setup.stamp_top_mm() == 24.0
    body_top = setup.body_top_mm(3)
    assert body_top > 20 + 6
    assert body_top < setup.height_mm - setup.margin_bottom_mm
    assert setup.body_top_mm(0) == 20.0


def test_rtf_editor_ha_pagina_margini_intestazione_e_interruzioni():
    content = build_rtf_export("", {"page_orientation": "orizzontale", "stamp_position": "top-right"}, timbro=_Timbro(), html_content=EDITOR_HTML)

    assert content.startswith("{\\rtf1")
    assert "\\paperw16838\\paperh11906" in content
    assert "\\landscape" in content and "\\lndscpsxn" in content
    assert "\\margl" in content and "\\margt" in content and "\\margb" in content
    assert "{\\header " in content and "STUDIO LEGALE ROSSI" in content and "\\qr" in content
    assert content.count("STUDIO LEGALE ROSSI") == 1
    assert "\\pagebb" in content
    assert "{\\strike barrato}" in content
    assert "\\qc" in content
    assert "Citt\\u224?" in content


def test_rtf_testo_semplice_e_caratteri_fuori_bmp():
    setup = page_setup_from_layout({})
    content = build_rtf_document(
        setup=setup,
        body_font="Times New Roman",
        stamp_font="Courier New",
        font_size_pt=12,
        line_height=1.5,
        text_align="justify",
        stamp_lines=[],
        text="Primo paragrafo\n\nSecondo {paragrafo}",
    )
    assert "Times New Roman" in content
    assert "Secondo \\{paragrafo\\}" in content
    assert "{\\header" not in content
    assert rtf_escape("😀") == "\\u-10179?\\u-8704?"


def test_docx_editor_ha_sezione_intestazione_e_interruzione_di_pagina():
    buffer = build_docx_export("", {"page_orientation": "orizzontale", "margin_left_mm": 30, "stamp_position": "top-center"}, timbro=_Timbro(), html_content=EDITOR_HTML)
    document = Document(io.BytesIO(buffer.getvalue()))
    section = document.sections[0]

    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert round(section.page_width.mm) == 297 and round(section.page_height.mm) == 210
    assert round(section.left_margin.mm) == 30
    assert "STUDIO LEGALE ROSSI" in "\n".join(paragraph.text for paragraph in section.header.paragraphs)
    body = [paragraph for paragraph in document.paragraphs if paragraph.text]
    assert "STUDIO LEGALE ROSSI" not in "\n".join(paragraph.text for paragraph in body)
    successiva = next(paragraph for paragraph in body if paragraph.text == "Pagina successiva")
    assert successiva.paragraph_format.page_break_before is True
    barrato = [run for paragraph in body for run in paragraph.runs if run.text == "barrato"]
    assert barrato and barrato[0].font.strike is True


def test_pdf_editor_rispetta_interruzione_orientamento_e_paragrafi():
    from pypdf import PdfReader

    from pct.editor import html_to_pdf

    html = "<h1>Atto</h1>" + "<p>Il sottoscritto avvocato &amp; collega espone i fatti.</p>" * 3 + EDITOR_HTML
    pdf = html_to_pdf(html, "Atto", layout={"stamp_anchor": "page_margin", "page_orientation": "orizzontale"}, studio_timbro=_Timbro())
    reader = PdfReader(io.BytesIO(pdf))

    assert len(reader.pages) == 2
    width, height = float(reader.pages[0].mediabox.width), float(reader.pages[0].mediabox.height)
    assert width > height
    first = reader.pages[0].extract_text()
    assert "STUDIO LEGALE ROSSI" in first
    assert "& collega" in first
    assert "Pagina successiva" in reader.pages[1].extract_text()
    assert "STUDIO LEGALE ROSSI" in reader.pages[1].extract_text()
