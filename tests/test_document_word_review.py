"""Guardrail della revisione Word dopo la prova materiale su 8080."""
from copy import deepcopy
import json

from docx import Document
from docx.enum.section import WD_SECTION_START
import pymupdf
import pytest

from web.services.document_word_review import apply_review


def block(text, page, identity):
    return {'id': identity, 'page': page, 'text': text, 'kind': 'paragrafo',
            'rows': [], 'tratti': [], 'format': {'corpo': 12, 'grassetto': False}}


def review_file(tmp_path, originals, updated):
    path = tmp_path / 'review.json'
    path.write_text(json.dumps({'originali': originals, 'blocchi': updated}), encoding='utf-8')
    return path


def source_pdf(texts):
    pdf = pymupdf.open()
    for text in texts:
        pdf.new_page().insert_text((72, 72), text)
    return pdf


def test_revision_is_scoped_to_page_and_preserves_table_and_other_runs(tmp_path):
    document = Document()
    first = document.add_paragraph()
    first.add_run('Titolo ').bold = True
    first.add_run('Societa').italic = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Tabella'
    table.cell(0, 1).text = 'Cella'
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph('Societa')
    originals = [block('Societa', 1, 'a'), block('Societa', 2, 'b')]
    updated = deepcopy(originals)
    updated[1]['text'] = 'Società'
    with source_pdf(['Societa', 'Societa']) as pdf:
        expected, aliases = apply_review(document, review_file(tmp_path, originals, updated), pdf)
    assert first.text == 'Titolo Societa'
    assert first.runs[0].bold is True and first.runs[1].italic is True
    assert document.paragraphs[-1].text == 'Società'
    assert table.cell(0, 0).text == 'Tabella'
    assert expected[1]['Societa'] == 1 and expected[2]['Società'] == 1
    assert aliases == {(1, 'Societa'): 'Società'}


def test_table_text_revision_retains_cells(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Cella'
    table.cell(0, 1).text = 'Destra'
    originals = [block('Cella', 1, 'a')]
    updated = deepcopy(originals)
    updated[0]['text'] = 'Cella corretta'
    with source_pdf(['Cella Destra']) as pdf:
        apply_review(document, review_file(tmp_path, originals, updated), pdf)
    assert [cell.text for cell in table.rows[0].cells] == ['Cella corretta', 'Destra']


def test_missing_revision_is_not_silently_ignored(tmp_path):
    document = Document()
    document.add_paragraph('Altro testo')
    originals = [block('Cella', 1, 'a')]
    updated = deepcopy(originals)
    updated[0]['text'] = 'Correzione'
    with source_pdf(['Cella']) as pdf, pytest.raises(ValueError, match='non trovato'):
        apply_review(document, review_file(tmp_path, originals, updated), pdf)


def test_grouped_lines_keep_mixed_styles_and_page_scope(tmp_path):
    document = Document()
    first = document.add_paragraph()
    first.add_run('Societa ')
    first.add_run('in grassetto').bold = True
    first.add_run(' e corsivo').italic = True
    second = document.add_paragraph('Accenti: à è é ì ò ù.')
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph('Societa in grassetto e corsivo')
    document.add_paragraph('Accenti: à è é ì ò ù.')
    originals = [block('Societa in grassetto e corsivo Accenti: à è é ì ò ù.', 1, 'a')]
    updated = deepcopy(originals)
    updated[0]['text'] = 'Società in grassetto e corsivo Accenti corretti: à è é ì ò ù.'
    with source_pdf([originals[0]['text'], originals[0]['text']]) as pdf:
        apply_review(document, review_file(tmp_path, originals, updated), pdf)
    assert first.text == 'Società in grassetto e corsivo'
    assert second.text == 'Accenti corretti: à è é ì ò ù.'
    assert next(run for run in first.runs if run.text == 'in grassetto').bold
    assert next(run for run in first.runs if run.text == ' e corsivo').italic
    assert document.paragraphs[-2].text == 'Societa in grassetto e corsivo'


def test_reset_to_source_keeps_each_original_run_style(tmp_path):
    from docx.shared import Pt, RGBColor
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Normale ').font.size = Pt(12)
    blue = paragraph.add_run('blu')
    blue.font.color.rgb = RGBColor.from_string('0000CC')
    blue.font.size = Pt(10)
    originals = [block(paragraph.text, 1, 'a')]
    updated = deepcopy(originals)
    updated[0]['format'].update(corpo=0, famiglia='', colore='')
    with source_pdf([paragraph.text]) as pdf:
        apply_review(document, review_file(tmp_path, originals, updated), pdf)
    assert [run.font.size.pt for run in paragraph.runs] == [12, 10]
    assert paragraph.runs[-1].font.color.rgb == RGBColor.from_string('0000CC')
