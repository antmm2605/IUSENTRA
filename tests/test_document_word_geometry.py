"""Guardrail: geometry must not guess repeated occurrences or worsen a line."""
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.shared import Pt
import pytest


@pytest.fixture
def geometry(monkeypatch):
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / 'web' / 'services'))
    from web.services import document_word_geometry
    return document_word_geometry


def test_repeated_inline_and_standalone_line_are_not_wrongly_paired(tmp_path, monkeypatch, geometry):
    document = Document()
    document.add_paragraph('Prefazione PEC: esempio dello studio')
    second = document.add_paragraph('PEC: esempio dello studio')
    second.paragraph_format.space_before = Pt(15)
    word = tmp_path / 'input.docx'
    document.save(word)
    initial = tmp_path / 'initial.pdf'
    monkeypatch.setattr(geometry, 'render_word', lambda *args: initial)
    rows = defaultdict(list, {'PEC: esempio dello studio': [(0, 100, 12), (1, 100, 12)]})
    monkeypatch.setattr(geometry, '_lines', lambda *args: rows)
    assert geometry.align_rendered_lines(tmp_path / 'source.pdf', word) == initial
    assert Document(word).paragraphs[1].paragraph_format.space_before.pt == 15


def test_worsened_candidate_restores_word_and_initial_render(tmp_path, monkeypatch, geometry):
    document = Document()
    paragraph = document.add_paragraph('Un capoverso con testo normale e grassetto')
    paragraph.paragraph_format.space_before = Pt(10)
    word = tmp_path / 'input.docx'
    document.save(word)
    original_bytes = word.read_bytes()
    source = tmp_path / 'source.pdf'
    initial = tmp_path / 'initial.pdf'
    candidate = tmp_path / 'candidate.pdf'
    key = paragraph.text
    calls = []
    def render(*args):
        calls.append(args)
        return initial if len(calls) == 1 else candidate
    def lines(path, *args):
        baseline = 10 if path == source else 15 if path == initial else 2
        return defaultdict(list, {key: [(0, baseline, 12)]})
    monkeypatch.setattr(geometry, 'render_word', render)
    monkeypatch.setattr(geometry, '_lines', lines)
    assert geometry.align_rendered_lines(source, word) == initial
    assert word.read_bytes() == original_bytes
    assert len(calls) == 3  # paragraph and physical-line proposals are rejected


def test_improved_candidate_keeps_all_inline_runs_on_same_baseline(tmp_path, monkeypatch, geometry):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Cliente ')
    paragraph.add_run('Esempio Controllato').bold = True
    paragraph.paragraph_format.space_before = Pt(10)
    word = tmp_path / 'input.docx'
    document.save(word)
    source, initial, candidate = [tmp_path / name for name in ('source.pdf', 'initial.pdf', 'candidate.pdf')]
    key = paragraph.text
    calls = []
    def render(*args):
        calls.append(args)
        return initial if len(calls) == 1 else candidate
    monkeypatch.setattr(geometry, 'render_word', render)
    monkeypatch.setattr(geometry, '_lines', lambda path, *args: defaultdict(list, {key: [(0, 15 if path == initial else 10, 12)]}))
    assert geometry.align_rendered_lines(source, word) == candidate
    reopened = Document(word).paragraphs[0]
    assert reopened.paragraph_format.space_before.pt == 5
    assert reopened.runs[1].bold and reopened.text == key
    assert 'w:position' not in reopened._p.xml


def test_rejected_step_is_not_reintroduced_by_next_accepted_step(tmp_path, monkeypatch, geometry):
    document = Document()
    keys = ['Primo capoverso controllato', 'Secondo capoverso controllato']
    for key in keys:
        document.add_paragraph(key).paragraph_format.space_before = Pt(10)
    word = tmp_path / 'input.docx'
    document.save(word)
    source, initial, rejected, accepted = [tmp_path / name for name in (
        'source.pdf', 'initial.pdf', 'rejected.pdf', 'accepted.pdf')]
    renders = iter([initial, rejected, accepted, accepted])
    monkeypatch.setattr(geometry, 'render_word', lambda *args: next(renders))
    baselines = {source: (10, 30), initial: (15, 35), rejected: (2, 35), accepted: (15, 30)}
    monkeypatch.setattr(geometry, '_lines', lambda path, *args: defaultdict(list, {
        key: [(0, value, 12)] for key, value in zip(keys, baselines[path])}))
    assert geometry.align_rendered_lines(source, word) == accepted
    paragraphs = Document(word).paragraphs
    assert paragraphs[0].paragraph_format.space_before.pt == 10
    assert paragraphs[1].paragraph_format.space_before.pt == 5


def test_physical_line_keeps_space_runs_and_moves_mixed_text_together(tmp_path, monkeypatch, geometry):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Parola').bold = True
    paragraph.add_run(' ')
    paragraph.add_run('seconda parte').italic = True
    word = tmp_path / 'input.docx'
    document.save(word)
    source, initial, candidate = [tmp_path / name for name in ('source.pdf', 'initial.pdf', 'candidate.pdf')]
    calls = []
    def render(*args):
        calls.append(args)
        return initial if len(calls) == 1 else candidate
    monkeypatch.setattr(geometry, 'render_word', render)
    monkeypatch.setattr(geometry, '_lines', lambda path, *args: defaultdict(list, {
        paragraph.text: [(0, 15 if path == initial else 10, 12)]}))
    assert geometry.align_rendered_lines(source, word) == candidate
    from docx.oxml.ns import qn
    runs = Document(word).paragraphs[0].runs
    assert ''.join(run.text for run in runs) == paragraph.text
    assert runs[0].bold and runs[2].italic
    assert [run._r.rPr.find(qn('w:position')).get(qn('w:val')) for run in runs] == ['10'] * 3


def test_leading_space_does_not_replace_visible_font_metrics(monkeypatch, geometry):
    from types import SimpleNamespace
    page = SimpleNamespace(number=0, get_text=lambda mode: {'blocks': [{'lines': [{'spans': [
        {'text': ' ', 'origin': (10, 40), 'size': 12},
        {'text': 'Testo visibile', 'origin': (13, 42), 'size': 11.52},
    ]}]}]})
    class PDF:
        def __enter__(self):
            return [page]
        def __exit__(self, *args):
            return False
    monkeypatch.setattr(geometry.pymupdf, 'open', lambda path: PDF())
    assert geometry._lines(Path('source.pdf'))['Testo visibile'] == [(0, 42, 11.52)]


def test_list_marker_moves_with_its_exact_text_line(tmp_path, monkeypatch, geometry):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('-\t')
    paragraph.add_run('Punto controllato della lista')
    word = tmp_path / 'input.docx'
    document.save(word)
    source, initial, candidate = [tmp_path / name for name in ('source.pdf', 'initial.pdf', 'candidate.pdf')]
    calls = []
    def render(*args):
        calls.append(args)
        return initial if len(calls) == 1 else candidate
    monkeypatch.setattr(geometry, 'render_word', render)
    monkeypatch.setattr(geometry, '_lines', lambda path, *args: defaultdict(list, {
        'Punto controllato della lista': [(0, 15 if path == initial else 10, 12)],
        '-': [(0, 15 if path == initial else 10, 12)],
    }))
    assert geometry.align_rendered_lines(source, word) == candidate
    from docx.oxml.ns import qn
    runs = Document(word).paragraphs[0].runs
    assert ''.join(run.text for run in runs) == paragraph.text
    assert [run._r.rPr.find(qn('w:position')).get(qn('w:val')) for run in runs] == ['10', '10']
