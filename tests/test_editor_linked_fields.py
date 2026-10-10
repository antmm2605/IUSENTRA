from types import SimpleNamespace

import pytest
from flask import Flask, g
from pct.clienti import Cliente, TipoCliente, Indirizzo, DocumentoIdentita
from pct.editor import html_to_docx
from pct.editor_linked_fields import linked_fields_catalog
from pct.editor_linked_fields import template_linked_html
from pct.documento_fedele.da_docx import converti_docx
from web.bootstrap.editor_linked_fields_routes import register_editor_linked_fields_routes


def sources():
    client = Cliente(id='C1', tipo=TipoCliente.PERSONA_FISICA, nome='Anna', cognome='Prova', data_nascita='1980-04-03',
                     indirizzo_residenza=Indirizzo(via='Via della Prova', civico='12', comune='Roma'),
                     documento=DocumentoIdentita(numero='CA12345AA', data_rilascio='2020-05-02'))
    matter = SimpleNamespace(id='F1', id_cliente='C1', numero='2026/001', titolo='Prova controllata', tribunale='Tribunale di Roma', valore_causa=1234.56)
    return client, matter


def test_catalogue_groups_and_values_use_native_sources():
    client, matter = sources()
    catalog = {field['id']: field for field in linked_fields_catalog(cliente=client, fascicolo=matter)}
    assert catalog['cliente.data_nascita']['value'] == '03/04/1980'
    assert catalog['cliente.indirizzo_residenza.via']['value'] == 'Via della Prova'
    assert catalog['cliente.documento.data_rilascio']['value'] == '02/05/2020'
    assert catalog['fascicolo.valore_causa']['value'] == '€ 1.234,56'
    assert not catalog['fascicolo.giudice']['available']
    assert catalog['cliente.documento.numero']['group'] == 'Documento d’identità'
    assert not any('note_riservate' in key for key in catalog)


def test_wrong_client_is_never_used():
    client, matter = sources()
    matter.id_cliente = 'ALTRO'
    with pytest.raises(ValueError):
        linked_fields_catalog(cliente=client, fascicolo=matter)


def test_studio_address_parts_are_configured_values_not_a_stamp_guess():
    client, matter = sources()
    catalog = {field['id']: field for field in linked_fields_catalog(
        cliente=client, fascicolo=matter,
        config={'studio': {'citta': 'Torino', 'provincia': 'TO', 'cap': '10121'}},
        studio_timbro={'cap_citta_provincia': '00100 Roma (RM)'})}
    assert catalog['studio.citta']['value'] == 'Torino'
    assert catalog['studio.provincia']['value'] == 'TO'
    assert catalog['studio.cap']['value'] == '10121'
    empty = {field['id']: field for field in linked_fields_catalog(
        cliente=client, fascicolo=matter, studio_timbro={'cap_citta_provincia': '00100 Roma (RM)'})}
    assert not empty['studio.provincia']['available']


def test_process_dates_are_distinct_and_never_inferred_from_upload():
    client, matter = sources()
    matter.data_apertura = '2025-04-06'
    matter.data_prima_udienza = '2026-03-12'
    matter.data_prossima_udienza = '2026-06-09'
    matter.data_notifica_citazione = '2025-04-22'
    catalog = {item['id']: item for item in linked_fields_catalog(cliente=client, fascicolo=matter)}
    assert catalog['fascicolo.data_apertura']['value'] == '06/04/2025'
    assert catalog['fascicolo.data_prima_udienza']['value'] == '12/03/2026'
    assert catalog['fascicolo.data_prossima_udienza']['value'] == '09/06/2026'
    assert catalog['fascicolo.data_notifica_citazione']['value'] == '22/04/2025'
    assert not catalog['fascicolo.data_chiusura']['available']


def test_model_rebinds_lawyer_and_stamp_to_current_studio():
    client, matter = sources()
    first = linked_fields_catalog(cliente=client, fascicolo=matter,
        config={'pec': {'indirizzo': 'alfa@example.test'}},
        studio_timbro={'professionista_nome': 'Avvocato Alfa', 'pec': 'alfa@example.test', 'indirizzo_riga': 'Via Prima 1'})
    second = linked_fields_catalog(cliente=client, fascicolo=matter,
        config={'pec': {'indirizzo': 'beta@example.test'}},
        studio_timbro={'professionista_nome': 'Avvocato Beta', 'pec': 'beta@example.test', 'indirizzo_riga': 'Via Seconda 2'})
    original = '<p><span data-iu-linked-field="studio_timbro.professionista_nome" style="font-weight:700;font-style:italic;color:#0055aa">Avvocato Alfa</span>, <span data-iu-linked-field="studio_timbro.pec">alfa@example.test</span></p>'
    model, _ = template_linked_html(original, context=first)
    filled, missing = template_linked_html(model, context=second)
    assert not missing and 'Avvocato Beta' in filled and 'beta@example.test' in filled
    assert 'Avvocato Alfa' not in filled and 'alfa@example.test' not in filled
    assert 'font-weight:700;font-style:italic;color:#0055aa' in filled


def test_act_place_date_and_pec_use_current_settings_not_old_model():
    from datetime import datetime
    from pct.formatting import DISPLAY_TIMEZONE, format_date_it
    client, matter = sources()
    fields = linked_fields_catalog(cliente=client, fascicolo=matter,
        config={'studio': {'citta': 'Roma'}, 'pec': {'indirizzo': 'corrente@example.test'}},
        studio_timbro={'pec': 'vecchia@example.test'})
    catalog = {field['id']: field for field in fields}
    assert catalog['studio.citta']['value'] == 'Roma'
    assert catalog['studio_timbro.pec']['value'] == 'corrente@example.test'
    assert catalog['studio_timbro.pec']['source'] == 'studio'
    today = format_date_it(datetime.now(DISPLAY_TIMEZONE).date())
    assert catalog['documento.data_atto']['value'] == today
    original = '<p><span data-iu-linked-field="studio.citta">Taurianova</span>, lì <span data-iu-linked-field="documento.data_atto">23.02.2026</span></p>'
    filled, missing = template_linked_html(original, context=fields)
    assert not missing and 'Taurianova' not in filled and '23.02.2026' not in filled
    assert 'Roma' in filled and today in filled
    absent = {field['id']: field for field in linked_fields_catalog(studio_timbro={'pec': 'vecchia@example.test'})}
    assert not absent['studio_timbro.pec']['available']


def test_complete_stamp_rows_rebind_native_layout_and_current_pec():
    first = linked_fields_catalog(config={'pec': {'indirizzo': 'prima@example.test'}},
        studio_timbro={'studio_nome': 'Studio Prima', 'indirizzo_riga': 'Via Prima 1',
                       'cap_citta_provincia': '00100 Roma (RM)', 'telefono': '123', 'pec': 'obsoleta@example.test'})
    second = linked_fields_catalog(config={'pec': {'indirizzo': 'seconda@example.test'}},
        studio_timbro={'studio_nome': 'Studio Seconda', 'indirizzo_riga': 'Via Seconda 2',
                       'cap_citta_provincia': '20100 Milano (MI)', 'telefono': '456', 'pec': 'altra-obsoleta@example.test'})
    original = ''.join(f'<p style="text-align:center"><span data-iu-linked-field="studio_timbro.{key}">Vecchio</span></p>'
                       for key in ('riga_studio', 'riga_indirizzo', 'riga_recapiti'))
    filled, missing = template_linked_html(original, context=first)
    assert not missing and 'prima@example.test' in filled and 'obsoleta' not in filled
    rebound, missing = template_linked_html(filled, context=second)
    assert not missing and 'STUDIO SECONDA' in rebound and 'Via Seconda 2' in rebound
    assert 'seconda@example.test' in rebound and 'prima@example.test' not in rebound
    assert '00100' not in rebound and '20100' in rebound


def test_model_updates_pec_destination_and_removes_mailto_on_non_email_fields():
    fields = linked_fields_catalog(config={'pec': {'indirizzo': 'attuale@example.test'}, 'studio': {'citta': 'Roma'}})
    original = '<p><a href="mailto:vecchia@example.test"><span data-iu-linked-field="studio_timbro.pec">vecchia@example.test</span></a></p><p><a href="mailto:vecchia@example.test"><span data-iu-linked-field="studio.citta">Taurianova</span></a></p>'
    filled, missing = template_linked_html(original, context=fields)
    assert not missing and 'vecchia@example.test' not in filled
    assert filled.count('mailto:') == 1 and 'mailto:attuale@example.test' in filled
    assert 'Roma' in filled
    model, _ = template_linked_html(original)
    assert 'vecchia@example.test' not in model
    rebound, missing = template_linked_html(model, context=fields)
    assert not missing and 'mailto:attuale@example.test' in rebound


@pytest.mark.parametrize('name,expected', [('Avv. Anna Prova', 'Anna Prova'), ('Avvocato Bruno Prova', 'Bruno Prova'), ('Avvocata Anna Prova', 'Anna Prova'), ('Anna Prova', 'Anna Prova')])
def test_lawyer_name_binding_keeps_existing_title_without_duplication(name, expected):
    fields = linked_fields_catalog(studio_timbro={'professionista_nome': name})
    template = '<p>con l’avv. <span data-iu-linked-field="studio_timbro.nome_professionista">Nome precedente</span></p>'
    filled, missing = template_linked_html(template, context=fields)
    assert not missing and expected in filled and 'avv. Avv.' not in filled


def test_hearing_date_is_rebound_at_every_occurrence_without_using_template_dates():
    client, matter = sources()
    original = ('<p><span data-iu-linked-field="fascicolo.data_prossima_udienza">12/03/2026</span></p>'
                '<p><span data-iu-linked-field="fascicolo.data_prossima_udienza">24/03/2026</span></p>')
    model, _ = template_linked_html(original)
    assert '12/03/2026' not in model and '24/03/2026' not in model
    empty, missing = template_linked_html(model, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert missing == ['Data della prossima udienza']
    assert empty.count('[Data della prossima udienza]') == 2
    matter.data_prossima_udienza = '2026-10-14'
    filled, missing = template_linked_html(model, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert filled.count('14/10/2026') == 2 and not missing
    matter.data_prossima_udienza = '2026-11-03'
    updated, missing = template_linked_html(model, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert updated.count('03/11/2026') == 2 and '14/10/2026' not in updated and not missing


def test_linked_word_controls_round_trip_preserves_format_and_context(tmp_path):
    html = '<p>Prima <span data-iu-linked-field="cliente.nome" data-iu-linked-matter="F1" data-iu-linked-client="C1"><strong><em>Anna</em></strong></span> dopo.</p>'
    path = tmp_path / 'linked.docx'
    path.write_bytes(html_to_docx(html))
    converted = converti_docx(path).html
    assert 'data-iu-linked-field="cliente.nome"' in converted
    assert 'data-iu-linked-matter="F1"' in converted
    assert 'data-iu-linked-client="C1"' in converted
    assert 'Anna' in converted and 'Prima' in converted and 'dopo.' in converted
    assert 'font-weight:700' in converted or '<strong>' in converted
    assert 'font-style:italic' in converted or '<em>' in converted
    second = tmp_path / 'second.docx'
    second.write_bytes(html_to_docx(converted))
    assert converti_docx(second).html.count('data-iu-linked-field=') == 1


def test_arbitrary_linked_path_is_rejected():
    with pytest.raises(ValueError):
        html_to_docx('<p><span data-iu-linked-field="cliente.note_riservate">Privato</span></p>')


def test_linked_fields_share_paragraph_tab_cursor_on_reopening(tmp_path):
    from docx import Document
    from docx.shared import Pt
    import re

    path = tmp_path / 'tabs-linked.docx'
    path.write_bytes(html_to_docx('<p><span data-iu-linked-field="cliente.nome">Anna</span><span data-iu-linked-field="cliente.cognome">Prova</span></p>'))
    document = Document(path)
    paragraph = document.paragraphs[0]
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(120))
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(240))
    paragraph.insert_paragraph_before('Titolo')
    paragraph._p.insert(1, paragraph.add_run('\t')._r)
    paragraph.add_run('\tFine')
    document.save(path)
    first = converti_docx(path).html
    widths = re.findall(r'data-iu-word-tab="(\d+)"', first)
    assert widths == ['2400', '4800']
    assert first.count('data-iu-linked-field=') == 2
    assert first.count('data-iu-word-tab-mark=') == 2
    path.write_bytes(html_to_docx(first))
    second = converti_docx(path).html
    assert re.findall(r'data-iu-word-tab="(\d+)"', second) == widths
    assert second.count('data-iu-linked-field=') == 2


def test_template_removes_old_values_and_refills_only_proven_fields():
    original = '<p>Cliente <span data-iu-linked-field="cliente.nome" data-iu-linked-client="VECCHIO"><strong>Nome precedente</strong></span>. <span data-iu-linked-field="cliente.documento.numero">CAOLD</span></p>'
    model, _ = template_linked_html(original)
    assert 'Nome precedente' not in model and 'CAOLD' not in model and 'VECCHIO' not in model
    client, matter = sources()
    filled, missing = template_linked_html(model, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert 'Anna' in filled and 'CA12345AA' in filled and not missing
    client.documento.numero = ''
    filled, missing = template_linked_html(model, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert missing == ['Numero documento'] and '[Numero documento]' in filled


def test_stamp_keeps_all_declared_tab_stops_after_two_roundtrips(tmp_path):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn

    path = tmp_path / 'stamp.docx'
    document = Document()
    paragraph = document.add_paragraph('\tStudio\n\t\t\t\tAvvocato\n\t\tCodice fiscale')
    paragraph.paragraph_format.left_indent = Pt(196.9)
    positions = [4014, 4074, 4132, 4478, 4544]
    for position in positions:
        paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(position / 20))
    document.save(path)
    for _ in range(2):
        path.write_bytes(html_to_docx(converti_docx(path).html))
        restored = Document(path).paragraphs[0]
        assert [int(tab.get(qn('w:pos'))) for tab in restored._p.findall(qn('w:pPr') + '/' + qn('w:tabs') + '/' + qn('w:tab'))] == positions
        assert restored.text == paragraph.text


def test_personal_model_sql_survives_legacy_catalogue_sync_and_retry(tmp_path):
    from pct.template_atti_repository import GestioneTemplateRepository
    repo = GestioneTemplateRepository(str(tmp_path / 'modelli.db'))
    model = {'id': 'modello1', 'titolo': 'Prova riutilizzabile', 'corpo': '<p>[Nome]</p>', 'collezione': 'Editor personale'}
    assert repo.create_editor_template(model) == 'modello1'
    assert repo.create_editor_template(model) == 'modello1'
    with pytest.raises(ValueError):
        repo.create_editor_template({**model, 'corpo': '<p>Diverso</p>'})
    repo.synchronize_templates([{'id': 'builtin', 'titolo': 'Integrato', 'corpo': 'Standard', 'builtin': True}], export_json=False)
    assert repo.get_template('modello1')['corpo'] == '<p>[Nome]</p>'
    repo.rebuild_from_payload([], source_label='mirror_storico')
    assert repo.get_template('modello1')['corpo'] == '<p>[Nome]</p>'
    assert not list(tmp_path.glob('*.json'))


@pytest.mark.parametrize('position,index', [('left', 0), ('center', 1), ('right', 2)])
def test_stamp_alignment_and_position_survive_word_round_trip(tmp_path, position, index):
    cells = ['<td style="width:33.333%;border:0"><p><br></p></td>'] * 3
    cells[index] = '<td style="width:33.333%;border:0"><p style="text-align:center">STUDIO DI PROVA</p><p style="text-align:center">Avvocato controllato</p></td>'
    html = f'<section class="iu-doc-pagina"><div data-iu-word-region="header" data-iu-word-kind="default"><table data-iu-editor-stamp="{position}" style="width:100%;border:0"><tbody><tr>{"".join(cells)}</tr></tbody></table></div><p>Corpo del documento</p></section>'
    path = tmp_path / 'stamp.docx'
    path.write_bytes(html_to_docx(html))
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document = Document(path)
    table = document.sections[0].header.tables[0]
    assert 'STUDIO DI PROVA' in table.cell(0, index).text
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER for paragraph in table.cell(0, index).paragraphs if paragraph.text)
    assert all(not table.cell(0, other).text.strip() for other in range(3) if other != index)
    from docx.oxml.ns import qn
    assert all(border.get(qn('w:val')) == 'nil' for cell in table.rows[0].cells for border in cell._tc.findall('./' + qn('w:tcPr') + '/' + qn('w:tcBorders') + '/*'))
    assert all(cell._tc.find('./' + qn('w:tcPr') + '/' + qn('w:tcBorders')) is not None for cell in table.rows[0].cells)
    converted = converti_docx(path).html
    assert f'data-iu-editor-stamp="{position}"' in converted
    assert 'Corpo del documento' in converted


def test_api_requires_both_permissions_and_tenant_and_native_fk(monkeypatch):
    client, matter = sources()
    app = Flask(__name__)
    permissions = {'clienti.leggi', 'fascicoli.leggi'}
    state = {'missing': False}
    @app.before_request
    def context():
        g.utente_corrente = SimpleNamespace(ha_permesso=lambda permission: permission in permissions)
        g.tenant_context_missing = state['missing']
    register_editor_linked_fields_routes(app,
        get_process_dates=lambda matter: {},
        get_fascicoli=lambda: SimpleNamespace(_studio_db=object(), get=lambda identifier: matter if identifier == 'F1' else None),
        get_clienti=lambda: SimpleNamespace(_studio_db=object(), get=lambda identifier: client if identifier == 'C1' else None))
    browser = app.test_client()
    queued = []
    monkeypatch.setattr('web.services.archivio_letture_runtime.lettura_dopo_evento', lambda fid: queued.append(fid) or True)
    assert browser.post('/api/editor/F1/recupera-campi').status_code == 202
    assert queued == ['F1']
    assert browser.post('/api/editor/ALTRO/recupera-campi').status_code == 404
    result = browser.get('/api/editor/F1/campi-collegati')
    assert result.status_code == 200 and result.json['clientId'] == 'C1'
    assert browser.get('/api/editor/ALTRO/campi-collegati').status_code == 404
    permissions.remove('clienti.leggi')
    assert browser.post('/api/editor/F1/recupera-campi').status_code == 403
    assert browser.get('/api/editor/F1/campi-collegati').status_code == 403
    permissions.add('clienti.leggi')
    state['missing'] = True
    assert browser.post('/api/editor/F1/recupera-campi').status_code == 503
    assert browser.get('/api/editor/F1/campi-collegati').status_code == 503
    state['missing'] = False
    matter.id_cliente = 'NONPRESENTE'
    assert browser.get('/api/editor/F1/campi-collegati').status_code == 409

def test_manual_value_stays_explicit_in_word_and_is_removed_from_reusable_model(tmp_path):
    original = '<p><span data-iu-linked-field="cliente.nome" data-iu-linked-matter="F1" data-iu-linked-client="C1" data-iu-linked-manual="true" contenteditable="false"><strong>Valore manuale</strong></span></p>'
    path = tmp_path / 'manual.docx'
    path.write_bytes(html_to_docx(original))
    converted = converti_docx(path).html
    assert 'data-iu-linked-manual="true"' in converted
    assert 'Valore manuale' in converted
    reusable, _ = template_linked_html(converted)
    assert 'Valore manuale' not in reusable
    assert 'data-iu-linked-manual' not in reusable
    client, matter = sources()
    filled, missing = template_linked_html(reusable, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert 'Anna' in filled and not missing

@pytest.mark.parametrize('mode,expected', [('upper', 'ANNA'), ('lower', 'anna'), ('title', 'Anna')])
def test_linked_case_survives_word_and_new_model_client(tmp_path, mode, expected):
    source = f'<p><span data-iu-linked-field="cliente.nome" data-iu-text-case="{mode}"><strong>PRECEDENTE</strong></span></p>'
    path = tmp_path / 'case.docx'
    path.write_bytes(html_to_docx(source))
    restored = converti_docx(path).html
    assert f'data-iu-text-case="{mode}"' in restored
    client, matter = sources()
    filled, missing = template_linked_html(restored, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert expected in filled and 'PRECEDENTE' not in filled and not missing
    assert '<strong' in filled or 'font-weight' in filled


def test_linked_long_date_survives_word_and_recompilation(tmp_path):
    source = '<p>Deposito: <span data-iu-linked-field="cliente.data_nascita" data-iu-date-format="long"><strong>30 aprile 2026</strong></span>.</p>'
    path = tmp_path / 'date.docx'
    path.write_bytes(html_to_docx(source))
    restored = converti_docx(path).html
    assert 'data-iu-date-format="long"' in restored
    client, matter = sources()
    filled, missing = template_linked_html(restored, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert '3 aprile 1980' in filled and '30 aprile 2026' not in filled and not missing


@pytest.mark.parametrize('mode,expected', [('dots', '03.04.1980'), ('long-padded', '03 aprile 1980')])
def test_original_date_presentation_survives_save_and_new_client(tmp_path, mode, expected):
    source = f'<p style="text-align:left"><span data-iu-linked-field="cliente.data_nascita" data-iu-date-format="{mode}"><strong><em>30.04.2026</em></strong></span></p>'
    path = tmp_path / 'source-format.docx'
    path.write_bytes(html_to_docx(source))
    restored = converti_docx(path).html
    assert 'text-align:left' in restored
    assert f'data-iu-date-format="{mode}"' in restored
    client, matter = sources()
    filled, missing = template_linked_html(restored, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert expected in filled and '30.04.2026' not in filled and not missing
    assert '<strong' in filled and '<em' in filled


def test_fiscal_code_grouping_preserves_word_position_and_rebinds(tmp_path):
    source = '<p><span data-iu-linked-field="cliente.codice_fiscale" data-iu-value-format="cf-grouped"><span style="--iu-word-position:7;position:relative;top:-7pt"><strong>RSS MRA 80A01 H501U</strong></span></span></p>'
    for index in range(2):
        path = tmp_path / f'grouped-{index}.docx'
        path.write_bytes(html_to_docx(source))
        source = converti_docx(path).html
        assert 'data-iu-value-format="cf-grouped"' in source
        assert '--iu-word-position:7' in source
    client, matter = sources()
    client.codice_fiscale = 'VRDLGU85B02F205X'
    filled, missing = template_linked_html(source, context=linked_fields_catalog(cliente=client, fascicolo=matter))
    assert 'VRD LGU 85B02 F205X' in filled and 'RSS MRA' not in filled and not missing
    assert '--iu-word-position:7' in filled


def test_pec_field_inside_link_survives_save_reopen_and_new_settings(tmp_path):
    source = '<p>PEC: <a href="mailto:prima@studio.it"><span data-iu-linked-field="studio_timbro.pec"><strong><em>prima@studio.it</em></strong></span></a> fine.</p>'
    path = tmp_path / 'pec.docx'
    path.write_bytes(html_to_docx(source))
    restored = converti_docx(path).html
    assert 'prima@studio.it' in restored
    assert 'data-iu-linked-field="studio_timbro.pec"' in restored
    assert 'mailto:prima@studio.it' in restored
    filled, missing = template_linked_html(restored, context=[{'id': 'studio_timbro.pec', 'value': 'nuova@studio.it', 'available': True}])
    assert 'mailto:nuova@studio.it' in filled and 'prima@studio.it' not in filled and not missing
    assert 'fine.' in filled


def test_field_directly_in_blank_editor_keeps_native_binding(tmp_path):
    html = '<section class="iu-doc-pagina"><span data-iu-linked-field="cliente.nome_completo" data-iu-linked-matter="F1" data-iu-linked-client="C1">Anna Prova</span><div><span data-iu-linked-field="cliente.data_nascita" data-iu-linked-matter="F1" data-iu-linked-client="C1">15/01/1980</span></div></section>'
    path = tmp_path / 'direct-fields.docx'
    path.write_bytes(html_to_docx(html))
    converted = converti_docx(path).html
    assert converted.count('data-iu-linked-field=') == 2
    assert 'Anna Prova' in converted and '15/01/1980' in converted
    path.write_bytes(html_to_docx(converted))
    assert converti_docx(path).html.count('data-iu-linked-field=') == 2
