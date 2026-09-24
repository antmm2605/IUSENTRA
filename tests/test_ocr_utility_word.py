"""Utility «Da scansione a Word o PDF»: catalogo, formato e documento prodotto."""

from __future__ import annotations

import io

from docx import Document
from flask import Flask, g

from pct.editor import html_to_docx
from web.blueprints.api_v1_document_tools import api_v1_document_tools
from web.services.documento_testo_riconosciuto import html_consentito, pdf_da_testo
from web.services.react_strumenti_legali_bridge import build_react_strumenti_legali_payload


def _app() -> Flask:
    app = Flask(__name__)

    @app.before_request
    def _autentica():
        g.utente_corrente = object()

    app.register_blueprint(api_v1_document_tools, url_prefix="/api/v1/ui/document-tools")
    return app


def test_il_catalogo_dichiara_la_utility_con_componente_dedicato():
    payload = build_react_strumenti_legali_payload(
        catalogo=[{"id": "ocr_documento_word", "title": "Da scansione a Word", "categoria": "Utility", "componente": "ocr-documento"}, {"id": "altro", "title": "Altro", "categoria": "X", "componente": "non-ammesso"}],
        form_state={},
        opzioni={},
    )
    voci = {voce["id"]: voce for voce in payload["strumenti"]}
    assert voci["ocr_documento_word"]["componente"] == "ocr-documento" and voci["ocr_documento_word"]["reso_in_react"]
    assert voci["altro"]["componente"] == "" and not voci["altro"]["reso_in_react"]
    assert payload["categorie"] == ["Utility", "X"]


def test_l_endpoint_restituisce_word_o_pdf_secondo_il_formato():
    client = _app().test_client()
    html = '<h1>Atto</h1><ol type="a" start="2"><li>secondo</li></ol><p style="text-align:justify">Testo.</p>'
    word = client.post("/api/v1/ui/document-tools/documento-testo-riconosciuto", data={"html": html, "nome": "scansione.pdf", "formato": "docx"})
    assert word.status_code == 200 and word.mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "testo riconosciuto.docx" in word.headers["Content-Disposition"]
    pdf = client.post("/api/v1/ui/document-tools/documento-testo-riconosciuto", data={"html": html, "nome": "scansione.pdf", "formato": "pdf"})
    assert pdf.status_code == 200 and pdf.mimetype == "application/pdf" and pdf.data.startswith(b"%PDF-")
    assert "testo riconosciuto.pdf" in pdf.headers["Content-Disposition"]


def test_gli_elenchi_con_tipo_e_partenza_arrivano_in_word_con_il_loro_segno():
    html = html_consentito('<ol type="a" start="2"><li>secondo</li><li>terzo</li></ol><ol type="I"><li>Premessa</li></ol><ol><li>uno</li></ol><ol type="x" start="abc"><li>senza attributi validi</li></ol>')
    assert 'type="x"' not in html and 'start="abc"' not in html
    paragrafi = [voce.text for voce in Document(io.BytesIO(html_to_docx(html, "prova", None))).paragraphs]
    assert paragrafi[:3] == ["b)\tsecondo", "c)\tterzo", "I.\tPremessa"]
    assert paragrafi[3] == "uno"


def test_il_pdf_del_testo_corretto_ha_il_nome_di_lavoro():
    dati, nome = pdf_da_testo("<p>Testo corretto</p>", "verbale udienza.jpg")
    assert dati.startswith(b"%PDF-") and nome == "verbale udienza - testo riconosciuto.pdf"


def test_colore_carattere_e_corpo_del_capoverso_arrivano_in_word():
    """Il capoverso porta il proprio formato ai suoi run: prima Word lo perdeva."""
    html = html_consentito(
        '<h2 style="text-align:center;color:#1f57a4;font-family:\'Book Antiqua\';font-size:16pt">Tribunale di Bari</h2>'
        '<p style="text-align:justify;font-family:\'Times New Roman\';font-size:12pt">Il ricorrente <strong>Tizio</strong> espone</p>'
    )
    paragrafi = Document(io.BytesIO(html_to_docx(html, "prova", None))).paragraphs
    titolo = paragrafi[0].runs[0]
    assert titolo.font.name == "Book Antiqua" and titolo.font.size.pt == 16.0
    assert str(titolo.font.color.rgb) == "1F57A4"
    # anche il neretto dentro il capoverso resta nel carattere del capoverso
    assert [(run.font.name, run.font.size.pt) for run in paragrafi[1].runs] == [("Times New Roman", 12.0)] * 3


def test_della_pila_dell_importazione_fedele_word_riceve_la_famiglia_del_catalogo():
    """L'alias incorporato e il nome PostScript Word non li conosce; lo sfondo non e' il testo."""
    html = '<p style="font-family:\'iu-3fa2b1\', \'TimesNewRomanPS\', \'Times New Roman\', serif;font-size:16px;background-color:#ff0000">x</p>'
    run = Document(io.BytesIO(html_to_docx(html, "prova", None))).paragraphs[0].runs[0]
    assert run.font.name == "Times New Roman" and run.font.size.pt == 12.0
    assert run.font.color.type is None


def test_il_colore_del_capoverso_arriva_nel_pdf():
    import fitz  # type: ignore

    dati, _ = pdf_da_testo('<p style="color:#0000ff">pec@avvocati.it</p>', "verbale.jpg")
    documento = fitz.open(stream=dati, filetype="pdf")
    try:
        colori = {
            span["color"]
            for blocco in documento[0].get_text("dict")["blocks"]
            for riga in blocco.get("lines", [])
            for span in riga["spans"]
            if "pec" in span["text"]
        }
    finally:
        documento.close()
    assert colori == {0x0000FF}


def test_il_formato_dentro_la_riga_arriva_in_word_parola_per_parola():
    html = html_consentito(
        '<p>Il Tribunale <strong><u>rigetta</u></strong> e scrive a '
        '<span style="color:#0000ff">studio@pec.it</span> e <s>cancella</s>.</p>'
    )
    runs = {run.text: run for run in Document(io.BytesIO(html_to_docx(html, "prova", None))).paragraphs[0].runs}
    assert runs["rigetta"].bold and runs["rigetta"].underline
    assert str(runs["studio@pec.it"].font.color.rgb) == "0000FF"
    assert runs["cancella"].font.strike
    assert not runs["Il Tribunale "].bold and runs["Il Tribunale "].font.color.type is None
