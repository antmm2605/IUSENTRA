"""Risorse dei modelli SQL: solo font incorporabili, senza testo della fonte."""
from base64 import b64decode, b64encode
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import UUID
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED


def template_repository():
    from flask import g
    from pct.template_atti_repository import GestioneTemplateRepository, derive_template_repository_db_path
    from pct.postgres_runtime_support import database_config_to_dsn, resolve_runtime_postgres_dsn
    anchor = (getattr(g, 'data_paths', {}) or {}).get('TEMPLATE_ATTI_DB')
    if getattr(g, 'tenant_context_missing', False) or not anchor:
        raise ValueError('Contesto studio dei modelli non disponibile.')
    tenant = getattr(g, 'tenant', None)
    dsn = database_config_to_dsn(getattr(tenant, 'database', None)) if tenant is not None else resolve_runtime_postgres_dsn()
    return GestioneTemplateRepository(derive_template_repository_db_path(str(anchor)), postgres_dsn=dsn)


def shared_template_repository():
    """Catalogo di piattaforma: mai derivare il DSN o il percorso dal tenant."""
    from flask import current_app
    import os
    from pct.template_atti_repository import GestioneTemplateRepository
    root = current_app.config.get('PCT_DATA_ROOT')
    if not root:
        return None
    path = Path(root) / 'editor_shared_models' / 'template_repository.db'
    dsn = str(current_app.config.get('EDITOR_SHARED_MODELS_POSTGRES_DSN') or os.getenv('EDITOR_SHARED_MODELS_POSTGRES_DSN') or '')
    if not dsn and not path.is_file():
        return None
    return GestioneTemplateRepository(str(path), postgres_dsn=dsn, initialize_schema=False)


def resolve_editor_template(identifier, personal_repository, shared_repository=None):
    """Scope esplicito: nessuna ricerca in archivi di altri studi."""
    shared = identifier.startswith('shared:')
    key = str(UUID(identifier[7:] if shared else identifier))
    repo = shared_repository if shared else personal_repository
    if repo is None:
        return None, None
    item = repo.get_template(key)
    expected = 'Editor condiviso' if shared else 'Editor personale'
    if not item or item['collezione'] != expected or (shared and not item['builtin']):
        return None, None
    return item, repo


def font_only_package(source):
    if not source:
        return ''
    from docx import Document
    from web.services.document_word_fonts import prepare_word_source_fonts, embed_source_fonts
    with TemporaryDirectory(prefix='iusentra-model-fonts-') as directory:
        fonts, _ = prepare_word_source_fonts(source, Path(directory))
        if not fonts:
            return ''
        document = Document()
        for font in fonts:
            run = document.add_paragraph().add_run(chr(min(font['glyphs'])))
            run.font.name = font['family']
            run.bold, run.italic = font['bold'], font['italic']
        embed_source_fonts(document, fonts, deterministic=True)
        buffer = BytesIO()
        document.save(buffer)
        canonical = BytesIO()
        with ZipFile(buffer) as source_zip, ZipFile(canonical, 'w', compression=ZIP_DEFLATED) as output:
            for name in sorted(source_zip.namelist()):
                info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
                info.compress_type = ZIP_DEFLATED
                output.writestr(info, source_zip.read(name))
        return b64encode(canonical.getvalue()).decode('ascii')


def template_font_source(html, repository=None):
    from lxml import html as parser
    root = parser.fragment_fromstring(html or '<p/>', create_parent='div')
    identifiers = {node.get('data-iu-font-template') for node in root.iter() if node.get('data-iu-font-template')}
    if not identifiers:
        return None
    if len(identifiers) != 1:
        raise ValueError('Il documento contiene risorse di modelli discordanti.')
    identifier = identifiers.pop()
    item, _ = resolve_editor_template(identifier, repository or template_repository(), shared_template_repository() if identifier.startswith('shared:') else None)
    if not item:
        raise ValueError('Risorse del modello non disponibili nello studio.')
    blocks = [block for block in item['template_blocks'] if block['block_key'] == 'editor_fonts']
    if len(blocks) != 1 or len(blocks[0]['content']) > 70_000_000:
        raise ValueError('Caratteri del modello non disponibili. Il documento è preservato.')
    return b64decode(blocks[0]['content'], validate=True)


def mark_template_fonts(html, identifier, repository):
    item = repository.get_template(identifier[7:] if identifier.startswith('shared:') else identifier)
    if not any(block['block_key'] == 'editor_fonts' for block in item['template_blocks']):
        return html
    from lxml import html as parser
    root = parser.fragment_fromstring(html, create_parent='div')
    for node in root.iter():
        node.attrib.pop('data-iu-font-template', None)
    root[0].set('data-iu-font-template', identifier)
    return ''.join(parser.tostring(node, encoding='unicode') for node in root)
