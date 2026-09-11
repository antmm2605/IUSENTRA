"""Esportazione DOCX fedele all'editor atti.

Formato pagina, margini e orientamento dell'editor; timbro dello studio
nell'intestazione di sezione (ripetuto su ogni pagina); interruzioni di pagina,
titoli, elenchi e formattazione del testo (grassetto, corsivo, sottolineato,
barrato, allineamento per paragrafo).
"""

from __future__ import annotations

import re
from typing import Any, Mapping, Sequence

from pct.template_atti_page_setup import PageSetup, is_page_break_element

_STYLE_ALIGN = re.compile(r"text-align\s*:\s*(left|right|center|justify)", re.I)
_HEADING_SCALE = {"h1": 1.42, "h2": 1.16, "h3": 1.06, "h4": 1.0}


def _alignment_map() -> dict[str, Any]:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }


def apply_page_setup(document: Any, setup: PageSetup, *, stamp_line_count: int) -> None:
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm

    for section in document.sections:
        section.orientation = WD_ORIENT.LANDSCAPE if setup.landscape else WD_ORIENT.PORTRAIT
        section.page_width = Mm(setup.width_mm)
        section.page_height = Mm(setup.height_mm)
        section.left_margin = Mm(setup.margin_left_mm)
        section.right_margin = Mm(setup.margin_right_mm)
        section.bottom_margin = Mm(setup.margin_bottom_mm)
        section.top_margin = Mm(setup.body_top_mm(stamp_line_count))
        section.header_distance = Mm(setup.stamp_top_mm() if setup.stamp_at_top else setup.margin_top_mm)


def write_stamp_header(document: Any, setup: PageSetup, stamp_lines: Sequence[Mapping[str, Any]], *, font_name: str) -> None:
    if not stamp_lines:
        return
    from docx.shared import Pt

    alignment = _alignment_map()[setup.stamp_alignment]
    for section in document.sections:
        header = section.header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        paragraph.alignment = alignment
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = setup.stamp_line_height
        for index, line in enumerate(stamp_lines):
            run = paragraph.add_run(str(line.get("text") or ""))
            run.bold = bool(line.get("bold"))
            run.font.name = font_name
            run.font.size = Pt(setup.stamp_font_size_pt)
            if index < len(stamp_lines) - 1:
                run.add_break()


def append_html(
    document: Any,
    html: str,
    *,
    font_name: str,
    font_size: float,
    line_height: float,
    default_align: str,
) -> None:
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    from lxml import etree

    align_map = _alignment_map()
    default_alignment = align_map.get(default_align, align_map["justify"])
    try:
        root = etree.fromstring(f"<div>{html}</div>".encode("utf-8"), parser=etree.HTMLParser(encoding="utf-8"))
        body = root.find(".//body")
        container = body[0] if body is not None and len(body) else root
    except Exception:
        container = None
    if container is None:
        return

    state = {"page_break": False}

    def tag_of(node: Any) -> str:
        return str(getattr(node, "tag", "") or "").lower().split("}")[-1]

    def alignment_for(node: Any, fallback: Any) -> Any:
        match = _STYLE_ALIGN.search(str(node.get("style") or "")) if hasattr(node, "get") else None
        if not match and hasattr(node, "get") and node.get("align"):
            match = _STYLE_ALIGN.search(f"text-align:{node.get('align')}")
        return align_map.get(match.group(1).lower(), fallback) if match else fallback

    def add_run(paragraph: Any, text: str, marks: Mapping[str, bool], size: float) -> None:
        if not text:
            return
        run = paragraph.add_run(re.sub(r"\s*\n\s*", " ", text))
        run.bold = marks.get("bold") or None
        run.italic = marks.get("italic") or None
        run.underline = marks.get("underline") or None
        run.font.strike = marks.get("strike") or None
        run.font.superscript = marks.get("sup") or None
        run.font.subscript = marks.get("sub") or None
        run.font.name = font_name
        run.font.size = Pt(size)

    def marks_for(node: Any, marks: Mapping[str, bool]) -> dict[str, bool]:
        tag = tag_of(node)
        style = str(node.get("style") or "").lower().replace(" ", "") if hasattr(node, "get") else ""
        return {
            "bold": bool(marks.get("bold") or tag in {"strong", "b", "th"} or "font-weight:bold" in style or re.search(r"font-weight:[6-9]00", style)),
            "italic": bool(marks.get("italic") or tag in {"em", "i", "cite"} or "font-style:italic" in style),
            "underline": bool(marks.get("underline") or tag in {"u", "ins"} or "underline" in style),
            "strike": bool(marks.get("strike") or tag in {"s", "strike", "del"} or "line-through" in style),
            "sup": bool(marks.get("sup") or tag == "sup"),
            "sub": bool(marks.get("sub") or tag == "sub"),
        }

    def walk_inline(paragraph: Any, node: Any, marks: Mapping[str, bool], size: float) -> None:
        current = marks_for(node, marks)
        if node.text:
            add_run(paragraph, str(node.text), current, size)
        for child in list(node):
            child_tag = tag_of(child)
            if child_tag == "br":
                paragraph.add_run().add_break()
            elif child.get("data-iu-line-spacer") is not None:
                pass
            elif child_tag not in {"ul", "ol"}:
                walk_inline(paragraph, child, current, size)
            if child.tail:
                add_run(paragraph, str(child.tail), current, size)

    def new_paragraph(node: Any, *, style: str | None = None, fallback_alignment: Any = None, space_after: float | None = None) -> Any:
        paragraph = document.add_paragraph(style=style) if style else document.add_paragraph()
        paragraph.alignment = alignment_for(node, fallback_alignment if fallback_alignment is not None else default_alignment)
        paragraph.paragraph_format.line_spacing = line_height
        paragraph.paragraph_format.space_after = Pt(font_size * 0.82 if space_after is None else space_after)
        if state["page_break"]:
            paragraph.paragraph_format.page_break_before = True
            state["page_break"] = False
        return paragraph

    def process_list(node: Any, depth: int) -> None:
        ordered = tag_of(node) == "ol"
        base = "List Number" if ordered else "List Bullet"
        style = base if depth <= 1 else f"{base} {min(depth, 3)}"
        for item in node.findall("li"):
            try:
                paragraph = new_paragraph(item, style=style, space_after=font_size * 0.25)
            except KeyError:
                paragraph = new_paragraph(item, style=base, space_after=font_size * 0.25)
            walk_inline(paragraph, item, {}, font_size)
            for nested in item:
                if tag_of(nested) in {"ul", "ol"}:
                    process_list(nested, depth + 1)

    def process(node: Any) -> None:
        tag = tag_of(node)
        if is_page_break_element(tag, dict(node.attrib)):
            state["page_break"] = True
            return
        if tag in _HEADING_SCALE:
            paragraph = new_paragraph(node, fallback_alignment=align_map["center"] if tag in {"h1", "h2"} else default_alignment)
            paragraph.paragraph_format.space_before = Pt(font_size)
            walk_inline(paragraph, node, {"bold": True}, round(font_size * _HEADING_SCALE[tag], 1))
            return
        if tag in {"ul", "ol"}:
            process_list(node, 1)
            return
        if tag == "hr":
            new_paragraph(node, space_after=font_size * 0.4)
            return
        if tag == "table":
            for row in node.iter("tr"):
                paragraph = new_paragraph(row, fallback_alignment=align_map["left"], space_after=font_size * 0.3)
                for index, cell in enumerate([cell for cell in row if tag_of(cell) in {"td", "th"}]):
                    if index:
                        paragraph.add_run("\t")
                    walk_inline(paragraph, cell, {}, font_size)
            return
        if tag in {"p", "div", "section", "article", "blockquote", "pre", "address", "figure"}:
            block_children = [child for child in node if tag_of(child) in {"p", "div", "ul", "ol", "table", "h1", "h2", "h3", "h4", "hr", "blockquote"}]
            if block_children and not (node.text or "").strip():
                for child in node:
                    process(child)
                return
            paragraph = new_paragraph(node)
            if tag == "blockquote":
                from docx.shared import Mm

                paragraph.paragraph_format.left_indent = Mm(12)
                walk_inline(paragraph, node, {"italic": True}, font_size)
            else:
                walk_inline(paragraph, node, {}, font_size)
            return
        for child in list(node):
            process(child)

    for child in list(container):
        process(child)
    if state["page_break"]:
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def append_text(document: Any, text: str, *, font_name: str, font_size: float, line_height: float, default_align: str) -> None:
    from docx.shared import Pt

    alignment = _alignment_map().get(default_align, _alignment_map()["justify"])
    for block in str(text or "").splitlines():
        paragraph = document.add_paragraph()
        paragraph.alignment = alignment
        paragraph.paragraph_format.line_spacing = line_height
        run = paragraph.add_run(block)
        run.font.name = font_name
        run.font.size = Pt(font_size)
