"""Integrità dei riferimenti font degli elenchi nell'RTF locale."""

import pytest

from pct.editor_export import _verifica_font_elenchi_rtf


def test_regioni_rtf_riprendono_misure_e_spaziature_native():
    from io import BytesIO
    from docx import Document
    from docx.shared import Pt
    from pct.editor_export import _regioni_rtf

    doc = Document()
    doc.sections[0].footer_distance = Pt(28)
    p = doc.sections[0].footer.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    stream = BytesIO()
    doc.save(stream)
    rtf = (b'{\\rtf1\\sectd\\footery680\\titlepg'
           b'{\\footer\\pard\\plain\\s0\\sa200{Testo {\\field{\\*\\fldinst PAGE}{\\fldrslt 1}}}\\par}'
           b' Corpo \\{\\\\footery999\\}}')
    risultato = _regioni_rtf(rtf, stream.getvalue())
    assert b'\\footery560 ' in risultato
    assert b'\\titlepg' not in risultato
    assert b'\\sb0\\sa0\\sl240\\slmult1 {Testo' in risultato
    assert b'{\\field{\\*\\fldinst PAGE}{\\fldrslt 1}}' in risultato
    assert b' Corpo \\{\\\\footery999\\}' in risultato


def test_regioni_rtf_non_eliminano_prima_pagina_distinta():
    from io import BytesIO
    from docx import Document
    from pct.editor_export import _regioni_rtf

    doc = Document()
    doc.sections[0].different_first_page_header_footer = True
    stream = BytesIO()
    doc.save(stream)
    risultato = _regioni_rtf(b'{\\rtf1\\sectd\\titlepg Corpo}', stream.getvalue())
    assert b'\\titlepg' in risultato
    with pytest.raises(ValueError, match='sezioni RTF'):
        _regioni_rtf(b'{\\rtf1 Corpo}', stream.getvalue())


@pytest.mark.parametrize('segno', [b'8226', b'9675', b'9642'])
def test_riferimento_assente_del_punto_riusa_arial_dichiarato(segno):
    dati = (b'{\\rtf1{\\fonttbl{\\f2\\fcharset0 Arial;}}'
            b'{\\*\\listtable{\\list{\\listlevel\\levelnfc23'
            b'{\\leveltext \\u8226 ?;}\\f16}}} Testo \\f2 invariato}')
    dati = dati.replace(b'8226', segno)
    atteso = dati.replace(b'\\f16', b'\\f2')
    assert _verifica_font_elenchi_rtf(dati) == atteso


def test_font_dichiarato_anche_symbol_resta_byte_identico():
    dati = (b'{\\rtf1{\\fonttbl{\\f16\\fcharset2 Symbol;}}'
            b'{\\*\\listtable{\\list{\\listlevel\\levelnfc23'
            b'{\\leveltext \\u8226 ?;}\\f16}}} Testo}')
    assert _verifica_font_elenchi_rtf(dati) == dati


@pytest.mark.parametrize('font,marcatore', [(b'Courier', b'\\u8226 ?'), (b'Arial', b'x')])
def test_riferimento_sconosciuto_non_viene_inventato(font, marcatore):
    dati = (b'{\\rtf1{\\fonttbl{\\f2 ' + font + b';}}'
            b'{\\*\\listtable{\\list{\\listlevel\\levelnfc23'
            b'{\\leveltext ' + marcatore + b';}\\f16}}}}')
    with pytest.raises(ValueError, match='non dichiarato'):
        _verifica_font_elenchi_rtf(dati)


def test_graffe_escapate_non_chiudono_il_catalogo():
    dati = (b'{\\rtf1{\\fonttbl{\\f2 Arial;}{\\f3 Font\\{nome\\};}}'
            b'{\\*\\listtable{\\list{\\listlevel\\levelnfc23'
            b'{\\leveltext \\u8226 ?;}\\f16}}} Corpo\\{\\}}')
    assert _verifica_font_elenchi_rtf(dati) == dati.replace(b'\\f16', b'\\f2')


def test_font_marcatore_rtf_riscontrato_nella_sorgente_word():
    from io import BytesIO
    from docx import Document
    from docx.oxml.ns import qn
    from pct.editor_export import _verifica_font_elenchi_rtf
    d = Document()
    n = d.part.numbering_part.element
    lvl = next(x for x in n.iter(qn('w:lvl')) if x.find(qn('w:rPr')) is not None)
    lvl.find(qn('w:lvlText')).set(qn('w:val'), '-')
    lvl.find(qn('w:rPr')).find(qn('w:rFonts')).set(qn('w:ascii'), 'Times New Roman')
    out = BytesIO()
    d.save(out)
    raw = br"{\rtf1{\fonttbl{\f0 Times New Roman;}}{\*\listtable{\listlevel\levelnfc23{\leveltext\u45 ?;}\f99}}}"
    fixed = _verifica_font_elenchi_rtf(raw, out.getvalue())
    assert br'\f99' not in fixed
    assert br'\f0}' in fixed
    with pytest.raises(ValueError):
        _verifica_font_elenchi_rtf(raw)
