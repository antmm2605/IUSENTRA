"""Worker separato: le tarature del convertitore non modificano altri flussi."""
from __future__ import annotations

import sys
from pathlib import Path
from statistics import median


def physical_line_y(line) -> float:
    """Measure the writing baseline, unaffected by bold/italic glyph boxes."""
    origins = [char.origin[1] for span in line.spans
               for char in getattr(span, 'chars', []) if char.origin]
    return median(origins) if origins else line.bbox.y0


def normalize_nested_hyperlinks(document) -> int:
    """Restore valid Word hyperlink siblings, preserving the converter's runs.

    Some PDF underlines are emitted as a hyperlink inside a run. Word ignores
    that invalid subtree, including its text. Move it to paragraph level and
    carry the original font properties into every contained run.
    """
    from copy import deepcopy
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    changed = 0
    for parent in list(document._element.body.iter(qn('w:r'))):
        if parent.getparent().tag != qn('w:p') or not parent.findall(qn('w:hyperlink')):
            continue
        properties = parent.find(qn('w:rPr'))
        pending = None
        for child in list(parent):
            if child.tag == qn('w:rPr'):
                continue
            if child.tag != qn('w:hyperlink'):
                if pending is None:
                    pending = OxmlElement('w:r')
                    if properties is not None:
                        pending.append(deepcopy(properties))
                pending.append(deepcopy(child))
                continue
            if pending is not None:
                parent.addprevious(pending)
                pending = None
            for run in child.findall(qn('w:r')):
                own = run.find(qn('w:rPr'))
                if properties is not None:
                    merged = deepcopy(properties)
                    if own is not None:
                        for item in own:
                            for inherited in list(merged.findall(item.tag)):
                                merged.remove(inherited)
                            merged.append(deepcopy(item))
                        run.remove(own)
                    run.insert(0, merged)
            parent.addprevious(child)
            changed += 1
        if pending is not None:
            parent.addprevious(pending)
        parent.getparent().remove(parent)
    return changed


def convert(source: Path, destination: Path, review_file: Path | None = None) -> int:
    import pymupdf
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from pdf2docx import Converter
    from pdf2docx.shape.Path import R
    from pdf2docx.layout.Section import Section
    from pdf2docx.text.TextBlock import TextBlock
    from pdf2docx.common.docx import set_columns
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    import numpy as np

    with pymupdf.open(source) as pdf:
        for page in pdf:
            if not page.get_text().strip() and page.get_images():
                return 20  # Non presentare una pagina immagine come testo modificabile.
    original_strokes = R.to_strokes
    original_section = Section.make_docx
    original_block = TextBlock.make_docx

    def measured_section(self, doc):
        if len(self) > 1 and all(block.is_text_block for column in self for block in column.blocks):
            # Le colonne native sono riquadri di testo indipendenti. Ancorarle
            # alla pagina evita bilanciamento automatico e migrazione del piè.
            set_columns(doc.sections[-1], [self.bbox.width], 0)
            for column in self:
                for block in column.blocks:
                    paragraph = doc.add_paragraph()
                    block._iusentra_fixed_column = True
                    block.make_docx(paragraph)
                    paragraph.paragraph_format.left_indent = Pt(0)
                    paragraph.paragraph_format.right_indent = Pt(0)
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    frame = OxmlElement('w:framePr')
                    for key, value in {'x': round(block.bbox.x0 * 20), 'y': round(block.bbox.y0 * 20),
                                       'w': round((block.bbox.width + 2) * 20), 'hAnchor': 'page',
                                       'vAnchor': 'page', 'wrap': 'none', 'anchorLock': '1'}.items():
                        frame.set(qn('w:' + key), str(value))
                    paragraph._p.get_or_add_pPr().append(frame)
            return
        # Una colonna nuova è un'interruzione di colonna nella stessa sezione,
        # non una sezione NEW_COLUMN: Writer la tratta come una pagina nuova.
        set_columns(doc.sections[-1], [column.bbox.width for column in self], self.space)
        for index, column in enumerate(self):
            if index:
                boundary = doc.add_paragraph()
                boundary.paragraph_format.space_before = Pt(0)
                boundary.paragraph_format.space_after = Pt(0)
                boundary.paragraph_format.line_spacing = Pt(0.1)
                boundary.add_run().add_break(WD_BREAK.COLUMN)
            column.make_docx(doc)

    def measured_block(self, paragraph):
        if not getattr(self, '_iusentra_fixed_column', False):
            for line, following in zip(self.lines, self.lines[1:]):
                if abs(physical_line_y(line) - physical_line_y(following)) > 1:
                    line.line_break = 1
            result = original_block(self, paragraph)
            gaps = [physical_line_y(following) - physical_line_y(line) for line, following in zip(self.lines, self.lines[1:])
                    if physical_line_y(following) - physical_line_y(line) > 1]
            if gaps and max(gaps) - min(gaps) <= 0.5:
                paragraph.paragraph_format.line_spacing = Pt(sum(gaps) / len(gaps))
            return result
        # Gli a capo presenti nel PDF non sono prosa da riformattare: tenerli
        # evita che colonne, elenchi e righe tabellari diventino un unico testo.
        for line, following in zip(self.lines, self.lines[1:]):
            if abs(physical_line_y(line) - physical_line_y(following)) > 1:
                line.line_break = 1
        result = original_block(self, paragraph)
        gaps = [physical_line_y(following) - physical_line_y(line) for line, following in zip(self.lines, self.lines[1:])
                if physical_line_y(following) - physical_line_y(line) > 1]
        if gaps and max(gaps) - min(gaps) <= 0.5:
            # L'interlinea relativa dipende dalle metriche del motore Word;
            # il passo uniforme misurato nella fonte resta invece in punti.
            paragraph.paragraph_format.line_spacing = Pt(sum(gaps) / len(gaps))
        return result

    def measured_strokes(self, width, color):
        return [{**stroke, 'width': width} for stroke in original_strokes(self, width, color)]

    R.to_strokes = measured_strokes
    Section.make_docx = measured_section
    TextBlock.make_docx = measured_block
    converter = Converter(str(source))
    try:
        converter.convert(str(destination), multi_processing=False)
    finally:
        converter.close()
        R.to_strokes = original_strokes
        Section.make_docx = original_section
        TextBlock.make_docx = original_block
    document = Document(destination)
    normalize_nested_hyperlinks(document)
    from document_word_fonts import prepare_source_fonts, embed_source_fonts
    with pymupdf.open(source) as font_source:
        source_fonts, font_config = prepare_source_fonts(font_source, destination.parent / 'source-fonts')
    for table in document.tables:
        grid = table._tbl.tblGrid
        columns = len(grid.gridCol_lst)
        matrix, widths = [], []
        for row in table._tbl.tr_lst:
            offset = 0
            for cell in row.tc_lst:
                properties = cell.tcPr
                span = int(properties.gridSpan.val) if properties.gridSpan is not None else 1
                width = properties.tcW
                if width is not None and width.type == 'dxa' and offset + span <= columns:
                    equation = [0] * columns
                    equation[offset:offset + span] = [1] * span
                    matrix.append(equation)
                    widths.append(int(width.w))
                offset += span
        if matrix:
            values, _, rank, _ = np.linalg.lstsq(np.array(matrix), np.array(widths), rcond=None)
            if rank == columns and min(values) > 0 and max(abs(np.array(matrix) @ values - widths)) <= 2:
                for column, width in zip(grid.gridCol_lst, values):
                    column.set(qn('w:w'), str(round(float(width))))
                table_width = table._tbl.tblPr.find(qn('w:tblW'))
                table_width.set(qn('w:type'), 'dxa')
                table_width.set(qn('w:w'), str(round(sum(values))))
        margins = OxmlElement('w:tblCellMar')
        for edge in ('top', 'left', 'bottom', 'right', 'start', 'end'):
            element = OxmlElement('w:' + edge)
            element.set(qn('w:w'), '0')
            element.set(qn('w:type'), 'dxa')
            margins.append(element)
        table._tbl.tblPr.append(margins)
        for element in table._tbl.iter():
            for name, value in list(element.attrib.items()):
                if name in {qn('w:w'), qn('w:sz')}:
                    try:
                        element.set(name, str(round(float(value))))
                    except ValueError:
                        pass
                if name == qn('w:color') and value.startswith('#'):
                    element.set(name, value[1:])
    expected = None
    aliases = None
    if review_file:
        from document_word_review import apply_review
        try:
            with pymupdf.open(source) as pdf:
                expected, aliases = apply_review(document, review_file, pdf)
        except (ValueError, KeyError, TypeError, IndexError):
            return 23
    embed_source_fonts(document, source_fonts)
    document.save(destination)
    from document_word_geometry import align_rendered_lines
    import shutil
    rendered = align_rendered_lines(source, destination, aliases, font_config)
    from collections import Counter
    import unicodedata
    import re
    with pymupdf.open(source) as original, pymupdf.open(rendered) as result:
        if len(original) != len(result):
            return 21
        for before, after in zip(original, result):
            if abs(before.rect.width - after.rect.width) > 0.5 or abs(before.rect.height - after.rect.height) > 0.5:
                return 21
            def words(page):
                return Counter(re.findall(r'\S+', unicodedata.normalize('NFKC', page.get_text())))
            expected_words = expected[before.number + 1] if expected else words(before)
            if expected_words - words(after):
                return 22
    shutil.copyfile(rendered, destination.with_suffix('.pdf'))
    return 0


if __name__ == '__main__':
    sys.exit(convert(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]) if len(sys.argv) > 3 else None))
