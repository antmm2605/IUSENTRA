from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from flask import Flask

from lex.formatting.document_export import build_docx_bytes
from lex.orchestrator_http import document_response


@pytest.mark.parametrize('origin', ['http://127.0.0.1:8080', 'https://app.iusentra.it'])
def test_word_receipt_source_preserves_label_and_clickable_destination(origin):
    label = 'CONSEGNA_Documento [JQ95-L01] [Notifica_ID:abc].eml'
    answer = r'- [CONSEGNA_Documento \[JQ95-L01\] \[Notifica_ID:abc\].eml](/fascicoli/F1/documenti/D1/visualizza)'
    document = Document(BytesIO(build_docx_bytes(title='Fonti delle ricevute', question='',
        answer=answer, source_origin=origin)))
    links = document._element.xpath('.//w:hyperlink')
    assert len(links) == 1
    assert ''.join(links[0].xpath('.//w:t/text()')) == label
    relationships = [r.target_ref for r in document.part.rels.values() if r.reltype == RELATIONSHIP_TYPE.HYPERLINK]
    assert relationships == [origin + '/fascicoli/F1/documenti/D1/visualizza']
    assert document.paragraphs[-1].style.name == 'List Bullet'


def test_word_keeps_external_source_and_text_around_it():
    document = Document(BytesIO(build_docx_bytes(title='Fonti', question='',
        answer='Leggi [Articolo 327](https://example.test/art327?x=1&y=2) e confronta.')))
    assert ''.join(document.paragraphs[-1]._p.xpath('.//w:t/text()')) == 'Leggi Articolo 327 e confronta.'
    assert any(r.target_ref == 'https://example.test/art327?x=1&y=2' for r in document.part.rels.values())


@pytest.mark.parametrize('destination', ['javascript:alert(1)', '//example.test', '/\\example.test', 'https://user:secret@example.test'])
def test_word_does_not_activate_untrusted_destinations(destination):
    document = Document(BytesIO(build_docx_bytes(title='Fonti', question='',
        answer=f'[Fonte]({destination})', source_origin='https://app.iusentra.it')))
    assert not document._element.xpath('.//w:hyperlink')


def test_word_relative_destination_without_origin_is_explicit_text():
    document = Document(BytesIO(build_docx_bytes(title='Fonti', question='',
        answer=r'[Ricevuta \[2\].eml](/fascicoli/F1/documenti/D1/visualizza)')))
    assert not document._element.xpath('.//w:hyperlink')
    text = ''.join(document.paragraphs[-1]._p.xpath('.//w:t/text()'))
    assert text == 'Ricevuta [2].eml (/fascicoli/F1/documenti/D1/visualizza)'


def test_http_export_keeps_paragraphs_and_uses_request_origin():
    app = Flask(__name__)
    deps = SimpleNamespace(build_docx_bytes=build_docx_bytes,
        infer_export_title=lambda **kw: 'Fonti delle ricevute',
        build_export_filename=lambda *args: 'fonti.docx')
    answer = '# Riepilogo\n\n- [Ricevuta \\[2\\]](/fascicoli/F1/documenti/D1/visualizza)\n\nConclusione.'
    with app.test_request_context(base_url='http://127.0.0.1:8080'):
        response = document_response(SimpleNamespace(dependencies=deps), data={
            'answer': answer, 'source_origin': 'https://example.invalid'})
        response.direct_passthrough = False
        document = Document(BytesIO(response.get_data()))
    assert any(p.style.name == 'Heading 2' and p.text == 'Riepilogo' for p in document.paragraphs)
    assert any(p.style.name == 'List Bullet' for p in document.paragraphs)
    assert document.paragraphs[-1].text == 'Conclusione.'
    assert any(r.target_ref == 'http://127.0.0.1:8080/fascicoli/F1/documenti/D1/visualizza' for r in document.part.rels.values())
