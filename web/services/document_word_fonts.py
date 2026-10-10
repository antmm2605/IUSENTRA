"""Font della fonte usati soltanto nel documento e nel suo render privato."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import hashlib
import re
import uuid
from xml.sax.saxutils import escape


def _name(value):
    return re.sub(r'[^a-z0-9]', '', str(value).split('+')[-1].casefold())


def _restore_document_unicode(font, source, xref):
    """Ricompone il collegamento Unicode dai codici dichiarati nel PDF.

    Nessun riconoscimento approssimativo: i glifi, le metriche e i diritti
    restano quelli incorporati nella fonte. La tabella serve solo al Word.
    """
    from fontTools.ttLib.tables._c_m_a_p import CmapSubtable
    kind, reference = source.xref_get_key(xref, 'ToUnicode')
    if kind != 'xref' or not re.fullmatch(r'\d+ \d+ R', reference):
        return False
    data = source.xref_stream(int(reference.split()[0]))
    if not data or len(data) > 2_000_000:
        return False
    text = re.sub(r'%[^\r\n]*', '', data.decode('ascii', errors='strict'))
    mappings = {}
    for block in re.findall(r'beginbfchar\s*(.*?)\s*endbfchar', text, re.S):
        for code, unicode in re.findall(r'<([0-9a-fA-F]{2,8})>\s*<([0-9a-fA-F]{4,8})>', block):
            character = bytes.fromhex(unicode).decode('utf-16-be')
            if len(character) == 1 and ord(character) >= 32:
                mappings[int(code, 16)] = ord(character)
    if not mappings or len(mappings) > 50000:
        return False
    encoded = [table.cmap for table in font['cmap'].tables if not table.isUnicode() and table.cmap]
    if len(encoded) != 1:
        return False
    unicode_map = {}
    for code, character in mappings.items():
        glyph = encoded[0].get(code)
        if glyph:
            if character in unicode_map and unicode_map[character] != glyph:
                return False
            unicode_map[character] = glyph
    if not unicode_map:
        return False
    table = CmapSubtable.newSubtable(12)
    table.platformID, table.platEncID, table.language = 3, 10, 0
    table.cmap = unicode_map
    font['cmap'].tables.append(table)
    return True


def prepare_source_fonts(source, directory: Path):
    from fontTools.ttLib import TTFont, TTLibError

    directory.mkdir(mode=0o700, parents=True, exist_ok=True)
    result, seen, total = [], set(), 0
    for page in source:
        for record in page.get_fonts(full=True):
            xref = record[0]
            if not xref or xref in seen:
                continue
            seen.add(xref)
            if len(seen) > 200:
                raise ValueError('Il documento contiene troppi caratteri incorporati')
            _, extension, _, data = source.extract_font(xref)
            if not data or extension not in {'ttf', 'otf'}:
                continue
            total += len(data)
            if len(data) > 10_000_000 or total > 50_000_000:
                raise ValueError('La libreria incorporata supera la dimensione consentita')
            try:
                with TTFont(BytesIO(data), lazy=False) as font:
                    if 'OS/2' not in font:
                        continue
                    rights = int(font['OS/2'].fsType)
                    # Editable/installable embedding only. Preview-only,
                    # restricted and bitmap-only fonts cannot enter editable Word.
                    if rights & 0x0200 or (rights & 0x000e and not rights & 0x0008):
                        continue
                    if rights & 0x0100 and '+' in record[3]:
                        continue
                    if not font.getBestCmap():
                        if not _restore_document_unicode(font, source, xref):
                            continue
                        restored = BytesIO()
                        font.save(restored)
                        data = restored.getvalue()
                    names = font['name']
                    family = names.getDebugName(16) or names.getDebugName(1)
                    if not family or len(family) > 100:
                        continue
                    style = names.getDebugName(17) or names.getDebugName(2) or ''
                    aliases = {_name(record[3]), _name(family), _name(names.getDebugName(4)), _name(names.getDebugName(6))}
                    bold = int(font['OS/2'].usWeightClass) >= 600
                    italic = bool(int(font['OS/2'].fsSelection) & 1) or 'italic' in style.casefold()
                    cmap = frozenset(font.getBestCmap())
            except (TTLibError, KeyError, ValueError, UnicodeError):
                continue
            path = directory / (hashlib.sha256(data).hexdigest() + '.' + extension)
            path.write_bytes(data)
            result.append({'family': family, 'aliases': aliases, 'bold': bold, 'italic': italic,
                           'data': data, 'glyphs': cmap})
    if result:
        config = directory / 'fonts.conf'
        config.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>'
                          '<include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                          '<dir>' + escape(str(directory.resolve())) + '</dir><cachedir>'
                          + escape(str((directory / 'cache').resolve())) + '</cachedir></fontconfig>', encoding='utf-8')
        return result, config
    return result, None


def _document_text_runs(document):
    """Corpo, intestazioni e piè di pagina effettivamente presenti nel package."""
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.text.run import Run

    roots = [document._element.body]
    roots.extend(relation.target_part.element for relation in document.part.rels.values()
                 if relation.reltype in {RT.HEADER, RT.FOOTER} and not relation.is_external)
    seen = set()
    for root in roots:
        if id(root) in seen:
            continue
        seen.add(id(root))
        for element in root.iter(qn('w:r')):
            yield Run(element, None)


def embed_source_fonts(document, fonts):
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    used = {}
    for run in _document_text_runs(document):
        name = _name(run.font.name)
        candidates = [font for font in fonts if name in font['aliases']]
        if not candidates:
            continue
        selected = next((font for font in candidates if font['bold'] == bool(run.bold) and font['italic'] == bool(run.italic)), candidates[0])
        if any(ord(character) not in selected['glyphs'] for character in run.text if not character.isspace()):
            continue
        run.font.name = selected['family']
        used[(selected['family'], selected['bold'], selected['italic'])] = selected
    if not used:
        return 0
    table = document.part.part_related_by(RT.FONT_TABLE)
    root = parse_xml(table.blob)
    for number, ((family, bold, italic), font) in enumerate(used.items()):
        entries = [entry for entry in root if entry.get(qn('w:name')) == family]
        entry = entries[0] if entries else OxmlElement('w:font')
        if not entries:
            entry.set(qn('w:name'), family)
            root.append(entry)
        tag = 'w:embed' + ('BoldItalic' if bold and italic else 'Bold' if bold else 'Italic' if italic else 'Regular')
        for previous in list(entry.findall(qn(tag))):
            entry.remove(previous)
        key = uuid.uuid4()
        mask = bytes.fromhex(key.hex)[::-1]
        data = bytearray(font['data'])
        for index in range(min(32, len(data))):
            data[index] ^= mask[index % 16]
        part = Part(PackURI('/word/fonts/iusentra-' + key.hex + '.odttf'),
                    'application/vnd.openxmlformats-officedocument.obfuscatedFont', bytes(data), document.part.package)
        relationship = table.relate_to(part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/font')
        embedded = OxmlElement(tag)
        embedded.set(qn('r:id'), relationship)
        embedded.set(qn('w:fontKey'), '{' + str(key).upper() + '}')
        entry.append(embedded)
    from lxml import etree
    table._blob = etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
    return len(used)


def prepare_word_source_fonts(data: bytes, directory: Path):
    """Riusa i font incorporati nel Word, senza installazioni globali."""
    from zipfile import ZipFile
    from lxml import etree
    from fontTools.ttLib import TTFont

    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
          'p': 'http://schemas.openxmlformats.org/package/2006/relationships'}
    directory.mkdir(mode=0o700, parents=True, exist_ok=True)
    fonts = []
    total = 0
    with ZipFile(BytesIO(data)) as archive:
        if 'word/_rels/fontTable.xml.rels' not in archive.namelist():
            return [], None
        if any(archive.getinfo(name).file_size > 2_000_000 for name in
               ('word/_rels/fontTable.xml.rels', 'word/fontTable.xml')):
            raise ValueError('La tabella dei caratteri incorporati è troppo grande.')
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        relations = etree.fromstring(archive.read('word/_rels/fontTable.xml.rels'), parser)
        targets = {item.get('Id'): item.get('Target') for item in relations
                   if item.get('TargetMode') != 'External'}
        table = etree.fromstring(archive.read('word/fontTable.xml'), parser)
        for entry in table:
            family = entry.get('{'+ns['w']+'}name')
            for tag in ('Regular', 'Bold', 'Italic', 'BoldItalic'):
                embedded = entry.find('w:embed'+tag, ns)
                if embedded is None:
                    continue
                target = targets.get(embedded.get('{'+ns['r']+'}id'), '')
                if not target.startswith('fonts/') or '..' in target.split('/') or '\\' in target:
                    raise ValueError('Percorso del carattere incorporato non valido.')
                info = archive.getinfo('word/'+target)
                total += info.file_size
                if len(fonts) >= 200 or info.file_size > 10_000_000 or total > 50_000_000:
                    raise ValueError('I caratteri incorporati superano la dimensione consentita.')
                key = uuid.UUID(embedded.get('{'+ns['w']+'}fontKey', ''))
                mask = key.bytes[::-1]
                restored = bytearray(archive.read(info))
                for index in range(min(32, len(restored))):
                    restored[index] ^= mask[index % 16]
                raw = bytes(restored)
                with TTFont(BytesIO(raw), lazy=False) as font:
                    rights = int(font['OS/2'].fsType)
                    if rights & 0x0200 or (rights & 0x000e and not rights & 0x0008):
                        continue
                    cmap = font.getBestCmap()
                    if not cmap:
                        raise ValueError('Il carattere incorporato non dichiara i glifi Unicode.')
                    names = font['name']
                    aliases = {_name(family), _name(names.getDebugName(1)), _name(names.getDebugName(4)), _name(names.getDebugName(6))}
                    fonts.append({'family': family, 'aliases': aliases,
                                  'bold': 'Bold' in tag, 'italic': 'Italic' in tag,
                                  'data': raw, 'glyphs': frozenset(cmap)})
                (directory / (hashlib.sha256(raw).hexdigest()+'.ttf')).write_bytes(raw)
    if not fonts:
        return [], None
    config = directory / 'fonts.conf'
    config.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>'
                      '<include ignore_missing="yes">/etc/fonts/fonts.conf</include><dir>'
                      + escape(str(directory.resolve())) + '</dir><cachedir>'
                      + escape(str((directory / 'cache').resolve())) + '</cachedir></fontconfig>', encoding='utf-8')
    return fonts, config


class SourceFontCoverageError(ValueError):
    """Il carattere della fonte non può rappresentare il testo aggiornato."""


def require_source_font_coverage(document, fonts):
    """Impedisce un export con glifi assenti dai font parziali della fonte."""
    incomplete = set()
    for run in _document_text_runs(document):
        candidates = [font for font in fonts if _name(run.font.name) in font['aliases']]
        if not candidates:
            continue
        selected = next((font for font in candidates if font['bold'] == bool(run.bold)
                         and font['italic'] == bool(run.italic)), candidates[0])
        if any(ord(character) not in selected['glyphs'] for character in run.text if not character.isspace()):
            incomplete.add(selected['family'])
    if incomplete:
        raise SourceFontCoverageError('Il font incorporato nella fonte non contiene tutte le lettere dei dati aggiornati. '
                         'Serve il font completo: ' + ', '.join(sorted(incomplete)) + '.')
