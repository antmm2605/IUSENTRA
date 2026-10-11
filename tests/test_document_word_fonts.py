"""Guardrail font incorporati dopo conversione reale e confronto delle due pagine."""
from io import BytesIO
import uuid
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
import pytest

from web.services.document_word_fonts import prepare_source_fonts, embed_source_fonts, prepare_word_source_fonts


def font_bytes(rights=8):
    builder = FontBuilder(1000, isTTF=True)
    builder.setupGlyphOrder(['.notdef', 'space', 'A'])
    builder.setupCharacterMap({32: 'space', 65: 'A'})
    glyphs = {}
    for name in ['.notdef', 'space', 'A']:
        pen = TTGlyphPen(None)
        if name == 'A':
            pen.moveTo((0, 0))
            pen.lineTo((250, 700))
            pen.lineTo((500, 0))
            pen.closePath()
        glyphs[name] = pen.glyph()
    builder.setupGlyf(glyphs)
    builder.setupHorizontalMetrics({name: (500, 0) for name in glyphs})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable({'familyName': 'QA Font', 'styleName': 'Regular',
                            'uniqueFontIdentifier': 'QA Font', 'fullName': 'QA Font', 'psName': 'QAFont'})
    builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200, fsType=rights)
    builder.setupPost()
    builder.setupMaxp()
    stream = BytesIO()
    builder.save(stream)
    return stream.getvalue()


class Source:
    def __init__(self, data):
        self.data = data

    def __iter__(self):
        return iter([self])

    def get_fonts(self, full=False):
        return [(10, 'ttf', 'TrueType', 'QAFont', 'F1', '')]

    def extract_font(self, xref):
        return 'QAFont', 'ttf', 'TrueType', self.data


def test_personal_model_fonts_preserve_bytes_without_source_text_and_survive_rebuild(tmp_path):
    from base64 import b64decode
    from pct.template_atti_repository import GestioneTemplateRepository
    from web.services.editor_template_fonts import font_only_package, template_font_source, mark_template_fonts
    fonts, _ = prepare_source_fonts(Source(font_bytes()), tmp_path / 'source')
    source = Document()
    run = source.add_paragraph().add_run('A')
    run.font.name = 'QA Font'
    embed_source_fonts(source, fonts)
    source.add_paragraph('Testo riservato della fonte, da non conservare nel modello')
    stream = BytesIO()
    source.save(stream)
    package = font_only_package(stream.getvalue())
    assert package == font_only_package(stream.getvalue())
    assert 'riservato' not in '\n'.join(p.text for p in Document(BytesIO(b64decode(package))).paragraphs)
    restored, _ = prepare_word_source_fonts(b64decode(package), tmp_path / 'restored')
    assert restored[0]['data'] == fonts[0]['data']
    repo = GestioneTemplateRepository(str(tmp_path / 'models.db'))
    payload = {'id': str(uuid.uuid4()), 'titolo': 'Modello con caratteri', 'corpo': '<p>A</p>', 'collezione': 'Editor personale'}
    identifier = repo.create_editor_template(payload, font_source=package)
    assert repo.create_editor_template(payload, font_source=package) == identifier
    with pytest.raises(ValueError):
        repo.create_editor_template(payload, font_source='diverso')
    html = mark_template_fonts(payload['corpo'], identifier, repo)
    assert template_font_source(html, repo) == b64decode(package)
    repo.rebuild_from_payload([])
    assert template_font_source(html, repo) == b64decode(package)
    other = GestioneTemplateRepository(str(tmp_path / 'other-tenant.db'))
    with pytest.raises(ValueError):
        template_font_source(html, other)


@pytest.mark.parametrize('rights', [2, 4, 0x0200])
def test_restricted_preview_or_bitmap_fonts_never_enter_editable_document(tmp_path, rights):
    fonts, config = prepare_source_fonts(Source(font_bytes(rights)), tmp_path / 'fonts')
    assert fonts == [] and config is None
    assert not list((tmp_path / 'fonts').glob('*.ttf'))


def test_embedded_font_keeps_original_bytes_and_document_relationship(tmp_path):
    data = font_bytes()
    fonts, config = prepare_source_fonts(Source(data), tmp_path / 'fonts')
    document = Document()
    run = document.add_paragraph().add_run('A')
    run.font.name = 'QAFont'
    assert embed_source_fonts(document, fonts) == 1
    assert run.font.name == 'QA Font'
    table = document.part.part_related_by(RT.FONT_TABLE)
    from docx.oxml import parse_xml
    root = parse_xml(table.blob)
    embedded = root.find('.//' + qn('w:embedRegular'))
    key = uuid.UUID(embedded.get(qn('w:fontKey')))
    blob = bytearray(table.related_parts[embedded.get(qn('r:id'))].blob)
    mask = bytes.fromhex(key.hex)[::-1]
    for index in range(32):
        blob[index] ^= mask[index % 16]
    assert bytes(blob) == data
    assert config.is_file()
    path = tmp_path / 'result.docx'
    document.save(path)
    with zipfile.ZipFile(path) as archive:
        assert len([name for name in archive.namelist() if name.endswith('.odttf')]) == 1


def test_incomplete_font_does_not_claim_new_character_coverage(tmp_path):
    fonts, _ = prepare_source_fonts(Source(font_bytes()), tmp_path / 'fonts')
    document = Document()
    run = document.add_paragraph().add_run('AB')
    run.font.name = 'QAFont'
    assert embed_source_fonts(document, fonts) == 0
    from web.services.document_word_fonts import require_source_font_coverage
    with pytest.raises(ValueError, match='font completo'):
        require_source_font_coverage(document, fonts)


def test_word_font_survives_two_native_editable_roundtrips(tmp_path):
    fonts, _ = prepare_source_fonts(Source(font_bytes()), tmp_path / 'pdf-fonts')
    document = Document()
    document.add_paragraph().add_run('A').font.name = 'QAFont'
    assert embed_source_fonts(document, fonts) == 1
    source = BytesIO()
    document.save(source)
    for index in range(2):
        recovered, config = prepare_word_source_fonts(source.getvalue(), tmp_path / str(index))
        assert config.is_file() and len(recovered) == 1
        assert recovered[0]['data'] == fonts[0]['data']
        assert recovered[0]['glyphs'] == frozenset({32, 65})
        output = Document()
        output.add_paragraph().add_run('A').font.name = 'QA Font'
        assert embed_source_fonts(output, recovered) == 1
        source = BytesIO()
        output.save(source)


def test_source_font_coverage_also_checks_header_and_footer(tmp_path):
    from web.services.document_word_fonts import require_source_font_coverage, SourceFontCoverageError
    fonts, _ = prepare_source_fonts(Source(font_bytes()), tmp_path / 'fonts')
    document = Document()
    for region in (document.sections[0].header, document.sections[0].footer):
        run = region.paragraphs[0].add_run('A')
        run.font.name = 'QAFont'
    require_source_font_coverage(document, fonts)
    document.sections[0].footer.paragraphs[0].runs[0].text = 'AB'
    with pytest.raises(SourceFontCoverageError, match='font completo'):
        require_source_font_coverage(document, fonts)


def test_source_font_resolver_rejects_cross_matter_and_changed_source(tmp_path):
    from types import SimpleNamespace
    from web.services.editor_word_source import editor_word_source
    import hashlib
    root = tmp_path / 'documents'
    (root / 'F1').mkdir(parents=True)
    path = root / 'F1' / 'source.docx'
    doc = Document()
    doc.save(path)
    source = SimpleNamespace(percorso='F1/source.docx', hash_sha256=hashlib.sha256(path.read_bytes()).hexdigest(), versioni=[])
    repo = SimpleNamespace(documents_dir=root)
    assert editor_word_source(repo, 'F1', source, lambda raw: raw) is None
    with pytest.raises(ValueError, match='non appartiene'):
        editor_word_source(repo, 'F2', source, lambda raw: raw)
    path.write_bytes(b'Changed')
    with pytest.raises(ValueError, match='impronta'):
        editor_word_source(repo, 'F1', source, lambda raw: raw)


def test_normal_editor_docx_export_keeps_source_font(tmp_path, monkeypatch):
    from pct import editor_export
    fonts, _ = prepare_source_fonts(Source(font_bytes()), tmp_path / 'fonts')
    source = Document()
    source.add_paragraph().add_run('A').font.name = 'QAFont'
    embed_source_fonts(source, fonts)
    raw = BytesIO()
    source.save(raw)
    target = Document()
    target.add_paragraph().add_run('A').font.name = 'QA Font'
    output = BytesIO()
    target.save(output)
    monkeypatch.setattr(editor_export, 'html_to_docx', lambda *args, **kwargs: output.getvalue())
    result = editor_export.esporta_documento_editor('<p>A</p>', formato='docx', fonte_word=raw.getvalue())
    recovered, _ = prepare_word_source_fonts(result, tmp_path / 'roundtrip')
    assert len(recovered) == 1 and recovered[0]['data'] == fonts[0]['data']
    assert Document(BytesIO(result)).paragraphs[0].text == 'A'


def test_legacy_tab_repair_requires_same_geometry_and_structure():
    from docx.shared import Pt
    from web.services.editor_word_source import restore_legacy_tab_stops
    source = Document()
    original = source.add_paragraph('\tStudio\n\t\tAvvocato')
    original.paragraph_format.left_indent = Pt(196.9)
    for position in (200.7, 203.7, 206.6):
        original.paragraph_format.tab_stops.add_tab_stop(Pt(position))
    target = Document()
    current = target.add_paragraph('\tStudio nuovo\n\t\tAvvocato nuovo')
    current.paragraph_format.left_indent = Pt(196.9)
    current.paragraph_format.tab_stops.add_tab_stop(Pt(200.7))
    assert restore_legacy_tab_stops(target, source, '<p>Studio</p>') == 1
    assert len(current.paragraph_format.tab_stops) == 3
    assert current.text == '\tStudio nuovo\n\t\tAvvocato nuovo'
    current.paragraph_format.tab_stops.clear_all()
    current.paragraph_format.tab_stops.add_tab_stop(Pt(200.7))
    current.paragraph_format.left_indent = Pt(120)
    assert restore_legacy_tab_stops(target, source, '<p>Studio</p>') == 0
    current.paragraph_format.left_indent = Pt(196.9)
    assert restore_legacy_tab_stops(target, source, '<p style="--iu-word-tab-stops:4014,left,none">Studio</p>') == 0


def test_shared_model_fonts_resolve_explicit_scope_and_remain_private_when_saved(tmp_path, monkeypatch):
    from base64 import b64decode
    from pct.template_atti_repository import GestioneTemplateRepository
    from web.services import editor_template_fonts as service
    shared = GestioneTemplateRepository(str(tmp_path / 'shared.db'))
    personal = GestioneTemplateRepository(str(tmp_path / 'tenant.db'))
    fonts, _ = prepare_source_fonts(Source(font_bytes()), tmp_path / 'source')
    document = Document()
    run = document.add_paragraph().add_run('A')
    run.font.name = 'QA Font'
    embed_source_fonts(document, fonts)
    buffer = BytesIO()
    document.save(buffer)
    package = service.font_only_package(buffer.getvalue())
    identifier = str(uuid.uuid4())
    payload = {'id': identifier, 'titolo': 'Comune', 'corpo': '<p>A</p>', 'collezione': 'Editor condiviso', 'builtin': True}
    shared.create_shared_editor_template(payload, font_source=package)
    monkeypatch.setattr(service, 'shared_template_repository', lambda: shared)
    compiled = service.mark_template_fonts(payload['corpo'], 'shared:' + identifier, shared)
    assert service.template_font_source(compiled, personal) == b64decode(package)
    with pytest.raises(ValueError):
        service.template_font_source(compiled.replace('shared:', ''), personal)
    own_id = str(uuid.uuid4())
    personal.create_editor_template({**payload, 'id': own_id, 'collezione': 'Editor personale', 'builtin': False}, font_source=package)
    private_html = service.mark_template_fonts(compiled, own_id, personal)
    assert 'shared:' not in private_html
    assert service.template_font_source(private_html, personal) == b64decode(package)
    other = GestioneTemplateRepository(str(tmp_path / 'other.db'))
    with pytest.raises(ValueError):
        service.template_font_source(private_html, other)
