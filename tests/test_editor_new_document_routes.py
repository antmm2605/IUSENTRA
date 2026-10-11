"""Creazione tramite repository SQL nativo; i test non sostituiscono la prova UI."""
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document
from flask import Flask, g
from pct.fascicoli import GestioneFascicoli, TipoFascicolo
from pct.storage import StudioDB
from web.bootstrap.editor_new_document_routes import register_editor_new_document_routes
from web.bootstrap.editor_finalize_routes import register_editor_finalize_routes

@pytest.fixture
def environment(tmp_path):
    database = StudioDB(str(tmp_path / "studio.db"))
    anchor = tmp_path / "fascicoli.json"
    def repository():
        return GestioneFascicoli(str(anchor), studio_db=database, carica_tutto=False)
    matter = GestioneFascicoli(str(anchor), studio_db=database).nuovo("Prova editor controllata", TipoFascicolo.CIVILE)
    app = Flask(__name__)
    app.config.update(TESTING=True, FASCICOLI_DB=str(anchor))
    permissions = {"fascicoli.leggi", "fascicoli.scrivi"}
    @app.before_request
    def context():
        g.utente_corrente = SimpleNamespace(username="operatore", ha_permesso=lambda p: p in permissions)
        g.data_paths = {"FASCICOLI_DB": str(anchor)}
    audits = []
    register_editor_new_document_routes(app, get_fascicoli=repository, audit=lambda *a, **kw: audits.append((a, kw)), encrypt_doc=lambda content: content)
    register_editor_finalize_routes(app, get_fascicoli=repository, audit=lambda *a, **kw: audits.append((a, kw)), encrypt_doc=lambda content: content, get_timbro=lambda: None)
    yield app.test_client(), repository, matter.id, permissions, audits
    database.chiudi()

def body():
    return {"titolo": "Documento controllato", "html": "<p>Bozza <strong>verificata</strong>.</p>", "command_id": str(uuid4())}

def test_native_sql_docx_round_trip_and_retry(environment):
    client, repo, matter, _, audits = environment
    request = body()
    first = client.post(f"/api/editor/{matter}/nuovo", json=request)
    assert first.status_code == 201
    repeated = client.post(f"/api/editor/{matter}/nuovo", json=request)
    assert repeated.json == first.json
    manager = repo()
    documents = manager.get(matter).documenti
    assert len(documents) == 1
    word = Document(BytesIO(manager.percorso_documento(matter, documents[0].id).read_bytes()))
    assert "Bozza verificata." in [p.text for p in word.paragraphs]
    assert any(run.bold for p in word.paragraphs for run in p.runs if run.text == "verificata")
    assert documents[0].fonte_documento == "REDAZIONE_STUDIO"
    assert len(audits) == 2

def test_reused_command_with_changed_content_is_refused(environment):
    client, repo, matter, _, _ = environment
    request = body()
    assert client.post(f"/api/editor/{matter}/nuovo", json=request).status_code == 201
    request["html"] = "<p>Diverso</p>"
    assert client.post(f"/api/editor/{matter}/nuovo", json=request).status_code == 409
    assert len(repo().get(matter).documenti) == 1

def test_separate_commands_create_separate_drafts(environment):
    client, repo, matter, _, _ = environment
    assert client.post(f"/api/editor/{matter}/nuovo", json=body()).status_code == 201
    assert client.post(f"/api/editor/{matter}/nuovo", json=body()).status_code == 201
    assert len(repo().get(matter).documenti) == 2

@pytest.mark.parametrize("permission", ["fascicoli.leggi", "fascicoli.scrivi"])
def test_permissions_refuse_creation(environment, permission):
    client, repo, matter, permissions, _ = environment
    permissions.remove(permission)
    assert client.post(f"/api/editor/{matter}/nuovo", json=body()).status_code == 403
    assert not repo().get(matter).documenti

def test_matters_are_sql_metadata(environment):
    client, _, matter, _, _ = environment
    result = client.get("/api/editor/nuovo/fascicoli")
    assert result.status_code == 200
    assert result.json["fascicoli"][0]["value"] == matter

def test_missing_matter_never_writes(environment):
    client, repo, matter, _, _ = environment
    assert client.post("/api/editor/INESISTENTE/nuovo", json=body()).status_code == 404
    assert not repo().get(matter).documenti


def test_preparation_link_is_persistent_and_second_session_reuses_document(environment, monkeypatch):
    from pct.controllo_studio.fonti import rif_fascicolo
    from web.services.preparazione_udienza_verbale import linked_document
    client, repo, matter, permissions, _ = environment
    permissions.update({'agenda.leggi', 'agenda.scrivi'})
    session = SimpleNamespace(id=str(uuid4()), id_fascicolo=matter, id_appuntamento='', verbale_preparazione='Bozza controllata.')
    monkeypatch.setattr('web.services.preparazione_udienza_verbale.source', lambda sid, fid: session)
    label = rif_fascicolo(repo().get(matter))['etichetta']
    payload = {'titolo': f'Preparazione verbale udienza {session.id}', 'command_id': str(uuid4()),
               'html': f'<h1>Preparazione del verbale di udienza</h1><p>{label}</p><p></p><p>Bozza controllata.</p>',
               'preparazione_id': session.id}
    first = client.post(f'/api/editor/{matter}/nuovo', json=payload)
    assert first.status_code == 201, first.json
    second = client.post(f'/api/editor/{matter}/nuovo', json={**payload, 'command_id': str(uuid4())})
    assert second.status_code == 201 and second.json == first.json
    documents = repo().get(matter).documenti
    assert len(documents) == 1
    assert linked_document(session.id, matter, documents)['href'].endswith('/' + first.json['document_id'] + '/editor')


def test_preparation_changed_source_cannot_create_unverified_document(environment, monkeypatch):
    client, repo, matter, permissions, _ = environment
    permissions.update({'agenda.leggi', 'agenda.scrivi'})
    session = SimpleNamespace(id=str(uuid4()), id_fascicolo=matter, id_appuntamento='', verbale_preparazione='Fonte salvata.')
    monkeypatch.setattr('web.services.preparazione_udienza_verbale.source', lambda sid, fid: session)
    result = client.post(f'/api/editor/{matter}/nuovo', json={**body(), 'preparazione_id': session.id})
    assert result.status_code == 409
    assert not repo().get(matter).documenti


def test_preparation_link_needs_agenda_permission(environment):
    client, repo, matter, _, _ = environment
    assert client.post(f'/api/editor/{matter}/nuovo', json={**body(), 'preparazione_id': str(uuid4())}).status_code == 403
    assert not repo().get(matter).documenti


@pytest.mark.parametrize('persistent,same_case,reference_issue', [(False, True, ''), (True, False, ''), (True, True, 'Riferimento discordante')])
def test_preparation_source_refuses_unadopted_or_other_case(monkeypatch, persistent, same_case, reference_issue):
    from web.services.preparazione_udienza_verbale import source
    sid = str(uuid4())
    session = SimpleNamespace(id=sid, id_fascicolo='CORRETTO' if same_case else 'ALTRO', reference_issue=reference_issue)
    manager = SimpleNamespace(write_protocol={'persistentCommands': persistent}, carica=lambda key: session)
    monkeypatch.setattr('web.helpers.get_wizard_pro', lambda: manager)
    with pytest.raises(ValueError):
        source(sid, 'CORRETTO')


def test_final_pdf_native_sql_retry_preserves_source(environment, monkeypatch):
    client, repo, matter, permissions, _ = environment
    request = body()
    source = client.post(f'/api/editor/{matter}/nuovo', json=request).json['document_id']
    manager = repo()
    manager.get(matter)
    original = manager.percorso_documento(matter, source).read_bytes()
    calls = []
    def convert(html, **kwargs):
        calls.append(html)
        return b'%PDF-1.7\nControlled converter guardrail'
    monkeypatch.setattr('pct.editor_export.esporta_documento_editor', convert)
    endpoint = f'/api/editor/{matter}/{source}/pdf-fascicolo'
    payload = {'html': request['html'], 'command_id': str(uuid4())}
    first = client.post(endpoint, json=payload)
    assert first.status_code == 201
    assert client.post(endpoint, json=payload).json == first.json
    assert len(calls) == 1 and len(repo().get(matter).documenti) == 2
    manager = repo()
    manager.get(matter)
    assert manager.percorso_documento(matter, source).read_bytes() == original
    assert client.post(endpoint, json={**payload, 'html': '<p>Diverso</p>'}).status_code == 409
    assert client.post(endpoint, json=['invalid']).status_code == 400
    assert client.post(endpoint, json={**payload, 'html': ''}).status_code == 400
    permissions.remove('fascicoli.scrivi')
    assert client.post(endpoint, json=payload).status_code == 403
