"""Word della revisione OCR con pagine e riquadri misurati, senza passaggi PDF."""
from __future__ import annotations

from copy import deepcopy
from collections import Counter
from io import BytesIO
from pathlib import Path
import tempfile
import re

from web.services.document_tools import DocumentToolError


def measured_word(html: str, title: str) -> bytes:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from lxml import etree
    from lxml.html import fragment_fromstring
    import pymupdf
    from pct.editor import html_to_docx
    from web.services.document_word_render import render_word

    root = fragment_fromstring(html, create_parent='div')
    pages = root.xpath('./section[@class="iu-doc-pagina"]')
    if not 1 <= len(pages) <= 100:
        raise DocumentToolError('La revisione Word accetta da 1 a 100 pagine per operazione.')
    document = Document()
    document.styles['Normal'].font.name = 'Times New Roman'
    document.styles['Normal'].font.size = Pt(12)
    expected = []
    for index, page in enumerate(pages):
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        width, height = float(page.get('data-larghezza', '595.28')), float(page.get('data-altezza', '841.89'))
        if not 72 <= width <= 2000 or not 72 <= height <= 2000:
            raise DocumentToolError('Le misure della pagina riconosciuta non sono valide.')
        section.page_width, section.page_height = Pt(width), Pt(height)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(0)
        page_text = []
        for block in page:
            if not isinstance(block.tag, str):
                continue
            part = Document(BytesIO(html_to_docx(etree.tostring(block, encoding='unicode'), title, None)))
            x, y = float(block.get('data-ocr-x', '0')), float(block.get('data-ocr-y', '0'))
            measured_width = float(block.get('data-ocr-width', str(width)))
            if not 0 <= x < width or not 0 <= y < height or measured_width <= 0:
                raise DocumentToolError('Un riquadro della revisione non rientra nella pagina.')
            # Il riquadro contiene testo Word vero; non è un'immagine dell'atto.
            # Un piccolo spazio laterale evita un a capo dovuto all'arrotondamento.
            frame_width = min(width - x, measured_width + 8)
            if block.get('data-ocr-estimated-font') == 'true':
                # Una scansione non dichiara il corpo. Lo si misura dal
                # riquadro; non si eredita il 12 pt del modello Word.
                for paragraph in part.paragraphs:
                    lines = paragraph.text.splitlines()
                    line_count = max(1, len(lines))
                    letter_height = float(block.get('data-ocr-height', '9')) / line_count
                    font = pymupdf.Font('tibo' if any(run.bold for run in paragraph.runs) else 'tiro')
                    unit_width = max((font.text_length(line, fontsize=1) for line in lines), default=1)
                    size = max(2, min(96, letter_height / 0.75, measured_width / max(unit_width, 0.1)))
                    for run in paragraph.runs:
                        run.font.size = Pt(round(size * 2) / 2)
            page_text.extend(''.join(element.text or '' for element in paragraph.iter(qn('w:t')))
                             for paragraph in part._element.body.iter(qn('w:p')))
            for element in part._element.body:
                if element.tag == qn('w:sectPr'):
                    continue
                element = deepcopy(element)
                if element.tag == qn('w:p'):
                    properties = element.get_or_add_pPr()
                    # Il livello riconosciuto non autorizza i colori e il
                    # corsivo del tema Word: il formato resta quello letto.
                    for old in list(properties.findall(qn('w:pStyle'))):
                        properties.remove(old)
                    for name in ('spacing', 'ind', 'keepNext', 'keepLines'):
                        for old in list(properties.findall(qn('w:' + name))):
                            properties.remove(old)
                    spacing = OxmlElement('w:spacing')
                    spacing.set(qn('w:before'), '0')
                    spacing.set(qn('w:after'), '0')
                    properties.append(spacing)
                    frame = OxmlElement('w:framePr')
                    for key, value in {'x': round(x * 20), 'y': round(y * 20), 'w': round(frame_width * 20), 'hAnchor': 'page', 'vAnchor': 'page', 'wrap': 'none', 'anchorLock': '1'}.items():
                        frame.set(qn('w:' + key), str(value))
                    properties.append(frame)
                elif element.tag == qn('w:tbl'):
                    properties = element.tblPr
                    position = OxmlElement('w:tblpPr')
                    for key, value in {'tblpX': round(x * 20), 'tblpY': round(y * 20), 'horzAnchor': 'page', 'vertAnchor': 'page'}.items():
                        position.set(qn('w:' + key), str(value))
                    properties.append(position)
                    for column in element.tblGrid.gridCol_lst:
                        column.set(qn('w:w'), str(round(frame_width * 20 / len(element.tblGrid.gridCol_lst))))
                document._element.body.insert(len(document._element.body) - 1, element)
        expected.append(Counter(re.findall(r'\S+', ' '.join(page_text))))
    output = BytesIO()
    document.save(output)
    with tempfile.TemporaryDirectory(prefix='iusentra-ocr-word-') as temporary:
        path = Path(temporary) / 'revisione.docx'
        path.write_bytes(output.getvalue())
        path.chmod(0o600)
        preview = render_word(path, Path(temporary) / 'render')
        with pymupdf.open(preview) as rendered:
            if len(rendered) != len(pages):
                raise DocumentToolError('Il Word della revisione ha cambiato il numero di pagine. La copia non è stata consegnata.')
            for index, page in enumerate(rendered):
                missing = expected[index] - Counter(re.findall(r'\S+', page.get_text()))
                if missing:
                    raise DocumentToolError('La resa Word perde parte del testo riconosciuto. La copia non è stata consegnata.')
                if any(x0 < -0.5 or y0 < -0.5 or x1 > page.rect.width + 0.5 or y1 > page.rect.height + 0.5
                       for x0, y0, x1, y1, *_ in page.get_text('words')):
                    raise DocumentToolError('Parte del testo Word esce dalla pagina. La copia non è stata consegnata.')
    return output.getvalue()
