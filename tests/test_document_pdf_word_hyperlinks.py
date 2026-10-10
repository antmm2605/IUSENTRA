"""PDF converter links must remain visible, ordered and formatted in Word."""
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from web.services.document_pdf_word_worker import normalize_nested_hyperlinks


def nested_link(parent, text, target, *, italic=False):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), target)
    run = OxmlElement('w:r')
    if italic:
        properties = OxmlElement('w:rPr')
        properties.append(OxmlElement('w:i'))
        run.append(properties)
    value = OxmlElement('w:t')
    value.text = text
    run.append(value)
    link.append(run)
    parent.append(link)
    return link


def test_nested_links_keep_mixed_text_order_styles_and_targets_after_reopen(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    outer = paragraph.add_run('Prima ')
    outer.bold = True
    outer.font.size = Pt(12)
    nested_link(outer._r, 'uno', 'rId10', italic=True)
    between = OxmlElement('w:t')
    between.text = ' e '
    outer._r.append(between)
    nested_link(outer._r, 'due', 'rId11')
    end = OxmlElement('w:t')
    end.text = ' dopo'
    outer._r.append(end)
    assert normalize_nested_hyperlinks(document) == 2
    assert normalize_nested_hyperlinks(document) == 0
    path = tmp_path / 'converted.docx'
    document.save(path)
    reopened = Document(path).paragraphs[0]._p
    assert ''.join(node.text or '' for node in reopened.iter(qn('w:t'))) == 'Prima uno e due dopo'
    assert [node.tag for node in reopened if node.tag != qn('w:pPr')] == [
        qn('w:r'), qn('w:hyperlink'), qn('w:r'), qn('w:hyperlink'), qn('w:r')]
    links = reopened.findall(qn('w:hyperlink'))
    assert [link.get(qn('r:id')) for link in links] == ['rId10', 'rId11']
    for link in links:
        properties = link.find(qn('w:r')).find(qn('w:rPr'))
        assert properties.find(qn('w:b')) is not None
        assert properties.find(qn('w:sz')).get(qn('w:val')) == '24'
    assert links[0].find(qn('w:r')).find(qn('w:rPr')).find(qn('w:i')) is not None


def test_valid_paragraph_link_is_unchanged_and_inner_font_overrides_outer():
    document = Document()
    paragraph = document.add_paragraph()
    valid = nested_link(paragraph._p, 'valido', 'rId1')
    original = valid.xml
    outer = paragraph.add_run()
    outer.bold = True
    nested = nested_link(outer._r, 'normale', 'rId2')
    run = nested.find(qn('w:r'))
    properties = OxmlElement('w:rPr')
    bold = OxmlElement('w:b')
    bold.set(qn('w:val'), '0')
    properties.append(bold)
    run.insert(0, properties)
    assert normalize_nested_hyperlinks(document) == 1
    assert valid.xml == original
    assert nested.find(qn('w:r')).find(qn('w:rPr')).find(qn('w:b')).get(qn('w:val')) == '0'
