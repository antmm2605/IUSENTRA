"""Il DOCX che rientra dall'editor deve ritrovare quello che aveva.

L'avvocato esporta l'atto in Word, lo modifica, lo ricarica. Se al rientro
perde gli allineamenti, il carattere o il colore, l'esportazione in Word non
e' un giro utile: e' un modo per rovinare l'atto.
"""

from __future__ import annotations

import re

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

from pct.documento_fedele.da_docx import DocxError, converti_docx


def _scrivi(costruisci, tmp_path, nome: str = "atto.docx"):
    documento = Document()
    costruisci(documento)
    percorso = tmp_path / nome
    documento.save(str(percorso))
    return converti_docx(percorso)


def _testo(html: str) -> str:
    return " ".join(re.sub(r"<[^>]+>", " ", html).split())


def test_gli_allineamenti_arrivano_tutti(tmp_path):
    def costruisci(d):
        titolo = d.add_paragraph("TRIBUNALE ORDINARIO DI NAPOLI")
        titolo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        corpo = d.add_paragraph("Con il presente atto si espone quanto segue.")
        corpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        firma = d.add_paragraph("Avv. Antonio Affinito")
        firma.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    html = _scrivi(costruisci, tmp_path).html
    assert "text-align:center" in html, "il titolo centrato e' finito a sinistra"
    assert "text-align:justify" in html, "il giustificato e' andato perso"
    assert "text-align:right" in html, "la firma a destra e' finita a sinistra"


def test_grassetto_corsivo_sottolineato_e_barrato_restano(tmp_path):
    def costruisci(d):
        p = d.add_paragraph()
        p.add_run("grassetto").bold = True
        p.add_run(" corsivo").italic = True
        p.add_run(" sottolineato").underline = True
        p.add_run(" barrato").font.strike = True

    html = _scrivi(costruisci, tmp_path).html
    assert "<strong>" in html and "<em>" in html
    assert "<u>" in html, "il sottolineato e' andato perso"
    assert "<s>" in html, "il barrato e' andato perso"


def test_il_colore_del_testo_arriva(tmp_path):
    def costruisci(d):
        p = d.add_paragraph()
        p.add_run("nero")
        rosso = p.add_run(" rilievo")
        rosso.font.color.rgb = RGBColor(0xBF, 0x19, 0x19)

    html = _scrivi(costruisci, tmp_path).html
    assert "#bf1919" in html.lower(), "il colore del testo e' andato perso"


def test_carattere_e_corpo_si_leggono_anche_quando_sono_ereditati(tmp_path):
    """In un DOCX quasi niente e' scritto dove lo si cerca: quasi tutto si eredita."""
    def costruisci(d):
        normale = d.styles["Normal"]
        normale.font.name = "Times New Roman"
        normale.font.size = Pt(12)
        d.add_paragraph("testo che non dichiara niente di suo")
        grande = d.add_paragraph()
        pezzo = grande.add_run("questo invece dichiara")
        pezzo.font.size = Pt(18)

    esito = _scrivi(costruisci, tmp_path)
    assert "Times New Roman" in esito.caratteri, (
        f"il carattere ereditato non e' stato letto: {esito.caratteri}"
    )
    assert "18" in esito.html, "il corpo dichiarato sul tratto e' andato perso"


def test_rientri_e_interlinea_arrivano(tmp_path):
    def costruisci(d):
        p = d.add_paragraph("paragrafo rientrato")
        p.paragraph_format.first_line_indent = Mm(10)
        p.paragraph_format.left_indent = Mm(15)
        p.paragraph_format.line_spacing = 1.5

    html = _scrivi(costruisci, tmp_path).html
    assert "text-indent" in html, "il rientro di prima riga e' andato perso"
    assert "margin-left" in html, "il rientro del paragrafo e' andato perso"
    assert "line-height" in html, "l'interlinea e' andata persa"


def test_gli_elenchi_restano_elenchi(tmp_path):
    def costruisci(d):
        d.add_paragraph("PREMESSO CHE")
        d.add_paragraph("primo punto", style="List Number")
        d.add_paragraph("secondo punto", style="List Number")
        d.add_paragraph("un punto puntato", style="List Bullet")

    html = _scrivi(costruisci, tmp_path).html
    assert "<ol>" in html, "l'elenco numerato e' diventato una sequenza di paragrafi"
    assert "<ul>" in html, "l'elenco puntato e' diventato una sequenza di paragrafi"
    assert html.count("<li>") == 3
    assert "primo punto" in _testo(html)


def test_la_tabella_con_intestazione_in_grassetto_ha_i_th(tmp_path):
    def costruisci(d):
        tabella = d.add_table(rows=3, cols=2)
        tabella.style = "Table Grid"
        dati = [("Voce", "Importo"), ("Contributo unificato", "EUR 518,00"),
                ("Compenso", "EUR 2.430,00")]
        for indice, (a, b) in enumerate(dati):
            tabella.rows[indice].cells[0].text = a
            tabella.rows[indice].cells[1].text = b
        for cella in tabella.rows[0].cells:
            for paragrafo in cella.paragraphs:
                for pezzo in paragrafo.runs:
                    pezzo.bold = True

    html = _scrivi(costruisci, tmp_path).html
    assert "<table" in html
    assert html.count("<th") == 2, "la riga di intestazione e' diventata una riga qualsiasi"
    assert html.count("<td") == 4
    assert "Contributo unificato" in _testo(html)


def test_una_tabella_in_mezzo_al_testo_resta_al_suo_posto(tmp_path):
    """Chi legge paragrafi e tabelle separatamente si ritrova le tabelle in fondo."""
    def costruisci(d):
        d.add_paragraph("PRIMA della tabella")
        tabella = d.add_table(rows=1, cols=2)
        tabella.rows[0].cells[0].text = "cella sinistra"
        tabella.rows[0].cells[1].text = "cella destra"
        d.add_paragraph("DOPO la tabella")

    html = _scrivi(costruisci, tmp_path).html
    posizione_prima = html.index("PRIMA")
    posizione_tabella = html.index("<table")
    posizione_dopo = html.index("DOPO")
    assert posizione_prima < posizione_tabella < posizione_dopo, (
        "la tabella non e' rimasta fra i due paragrafi"
    )


def test_formato_e_margini_della_pagina_sono_dichiarati(tmp_path):
    def costruisci(d):
        sezione = d.sections[0]
        sezione.page_width, sezione.page_height = Mm(210), Mm(297)
        sezione.left_margin = sezione.right_margin = Mm(25)
        d.add_paragraph("atto")

    esito = _scrivi(costruisci, tmp_path)
    assert esito.formato["formato"] == "A4"
    assert esito.formato["orientamento"] == "verticale"
    assert esito.formato["margini_pt"]["sinistro"] == pytest.approx(70.9, abs=1.0)


def test_una_pagina_orizzontale_viene_riconosciuta(tmp_path):
    def costruisci(d):
        sezione = d.sections[0]
        sezione.page_width, sezione.page_height = Mm(297), Mm(210)
        d.add_paragraph("prospetto")

    esito = _scrivi(costruisci, tmp_path)
    assert esito.formato["orientamento"] == "orizzontale"


def test_un_documento_vuoto_non_fa_saltare_la_conversione(tmp_path):
    esito = _scrivi(lambda d: None, tmp_path)
    assert esito.pagine
    assert esito.avvisi


def test_un_file_che_non_e_un_docx_viene_rifiutato(tmp_path):
    finto = tmp_path / "finto.docx"
    finto.write_bytes(b"non sono un documento")
    with pytest.raises(DocxError, match="illeggibile"):
        converti_docx(finto)
