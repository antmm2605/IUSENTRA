"""Preserva i segni OCR senza cambiare gli elenchi ordinari dell'editor."""
import io

import fitz
import pytest
from docx import Document

from web.services.documento_testo_riconosciuto import docx_da_testo, html_consentito, pdf_da_testo


@pytest.mark.parametrize("marker", ["5)", "3.", "(2)", "a)", "IV.", "1.2", "•", "-"])
def test_marker_ocr_rimane_identico_in_word_e_pdf(marker):
    html = f'<ol start="5"><li data-iu-ocr-marker="{marker}">Voce controllata</li></ol>'
    word, _ = docx_da_testo(html, "prova.pdf")
    assert Document(io.BytesIO(word)).paragraphs[0].text == marker + "\tVoce controllata"
    pdf, _ = pdf_da_testo(html, "prova.pdf")
    with fitz.open(stream=pdf, filetype="pdf") as document:
        assert marker in document[0].get_text()
        assert "Voce controllata" in document[0].get_text()


def test_numerazione_ordinaria_non_cambia():
    word, _ = docx_da_testo('<ol start="5"><li>Voce</li></ol>', "prova.pdf")
    assert Document(io.BytesIO(word)).paragraphs[0].text == "5.\tVoce"


@pytest.mark.parametrize("marker", ["<script>", "javascript:alert(1)", "x" * 1000, "a b", ""])
def test_marker_non_valido_non_entra_nel_documento(marker):
    sanitized = html_consentito(f'<ol><li data-iu-ocr-marker="{marker}">Voce</li></ol>')
    assert "data-iu-ocr-marker" not in sanitized
