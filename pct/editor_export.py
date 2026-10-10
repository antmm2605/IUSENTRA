"""Formati dell'editor ottenuti dalla stessa sorgente DOCX, solo in locale."""

from __future__ import annotations

import shutil
import re
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from pct.editor import html_to_docx


def _regioni_rtf(dati: bytes, sorgente: bytes) -> bytes:
    """Conserva le misure native che il convertitore perde dopo la prima sezione.

    Si correggono solo controlli esterni ai gruppi e spaziature esplicite dei
    paragrafi delle parti Word corrispondenti. Corpo, campi e immagini restano
    invariati; le sezioni devono corrispondere prima di qualsiasi modifica.
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(BytesIO(sorgente))
    tokens = []
    gruppi = []
    depth = 0
    bin_fine = 0
    for match in re.finditer(rb"\\(?:[a-zA-Z]+-?[0-9]* ?|[^a-zA-Z])|[{}]", dati):
        if match.start() < bin_fine:
            continue
        token = match.group()
        if token == b"{":
            depth += 1
            gruppi.append((match.start(), depth))
        elif token == b"}":
            depth -= 1
        else:
            tokens.append((match, depth))
            binary = re.fullmatch(rb"\\bin([0-9]+) ?", token)
            if binary:
                bin_fine = match.end() + int(binary.group(1))
    inizi = [m.start() for m, livello in tokens if livello == 1 and m.group().rstrip() == b"\\sectd"]
    if len(inizi) != len(doc.sections):
        raise ValueError("Le sezioni RTF non corrispondono alla sorgente Word.")
    cambi = []
    for indice, inizio in enumerate(inizi):
        fine = inizi[indice + 1] if indice + 1 < len(inizi) else len(dati)
        sezione = doc.sections[indice]
        for match, livello in tokens:
            if livello != 1 or not inizio <= match.start() < fine:
                continue
            token = match.group()
            for nome, misura in ((b"headery", sezione.header_distance), (b"footery", sezione.footer_distance)):
                if re.fullmatch(rb"\\" + nome + rb"-?[0-9]+ ?", token) and misura is not None:
                    cambi.append((match.start(), match.end(), b"\\" + nome + str(misura.twips).encode() + b" "))
            if token.rstrip() == b"\\titlepg" and not sezione.different_first_page_header_footer:
                cambi.append((match.start(), match.end(), b""))
        for match, livello in tokens:
            regione = re.fullmatch(rb"\\(header|footer)([lf]?) ?", match.group())
            if livello != 2 or not inizio <= match.start() < fine or not regione:
                continue
            start = match.start() - 1
            if dati[start:start + 1] != b"{":
                continue
            end = _fine_gruppo_rtf(dati, start)
            nome = regione.group(1).decode()
            prefisso = {b"": "", b"l": "even_page_", b"f": "first_page_"}[regione.group(2)]
            parte = getattr(sezione, prefisso + nome)
            # Tabelle richiedono una corrispondenza di celle: non si appiattiscono.
            if parte.tables:
                continue
            paragrafi = []
            par_inizio = None
            for controllo, livello_p in tokens:
                if livello_p == 2 and start < controllo.start() < end:
                    if controllo.group().rstrip() == b"\\pard":
                        par_inizio = controllo.start()
                    elif controllo.group().rstrip() == b"\\par" and par_inizio is not None:
                        paragrafi.append((par_inizio, controllo.start()))
                        par_inizio = None
            if len(paragrafi) != len(parte.paragraphs):
                continue
            for (par_inizio, par_fine), paragrafo in zip(paragrafi, parte.paragraphs):
                spacing = paragrafo._p.find("./" + qn("w:pPr") + "/" + qn("w:spacing"))
                if spacing is None:
                    continue
                controlli = b""
                for attr, comando in (("before", b"sb"), ("after", b"sa")):
                    valore = spacing.get(qn("w:" + attr))
                    if valore is not None:
                        controlli += b"\\" + comando + str(int(valore)).encode()
                line = spacing.get(qn("w:line"))
                rule = spacing.get(qn("w:lineRule"), "auto")
                if line is not None:
                    valore = int(line) * (-1 if rule == "exact" else 1)
                    controlli += b"\\sl" + str(valore).encode() + b"\\slmult" + (b"1" if rule == "auto" else b"0")
                # Il primo gruppo è il contenuto del paragrafo, dopo i suoi stili.
                gruppo = next((p for p, livello_g in gruppi if livello_g == 3 and par_inizio < p < par_fine), None)
                if controlli and gruppo is not None:
                    posizione = gruppo
                    cambi.append((posizione, posizione, controlli + b" "))
    for inizio, fine, valore in sorted(cambi, reverse=True):
        dati = dati[:inizio] + valore + dati[fine:]
    return dati


def _fine_gruppo_rtf(dati: bytes, inizio: int) -> int:
    profondita = 0
    indice = inizio
    while indice < len(dati):
        carattere = dati[indice]
        if carattere == 92:
            indice += 2
            continue
        if carattere == 123:
            profondita += 1
        elif carattere == 125:
            profondita -= 1
            if profondita == 0:
                return indice + 1
        indice += 1
    raise ValueError("Struttura RTF incompleta.")


def _verifica_font_elenchi_rtf(dati: bytes, sorgente: bytes | None = None) -> bytes:
    """Ripara i riferimenti dei tre segni Arial canonici degli elenchi.

    Il convertitore locale può scrivere nella listtable un indice font che
    manca dalla fonttbl. Il lettore sostituisce allora il punto e cambia
    l'altezza della riga. Si riusa la dichiarazione Arial già presente;
    riferimenti validi e ogni altro contenuto restano byte-identici.
    """
    font_inizio = dati.find(b"{\\fonttbl")
    liste_inizio = dati.find(b"{\\*\\listtable")
    if liste_inizio < 0:
        return dati
    if font_inizio < 0:
        raise ValueError("Catalogo caratteri RTF assente.")
    font_fine = _fine_gruppo_rtf(dati, font_inizio)
    liste_fine = _fine_gruppo_rtf(dati, liste_inizio)
    catalogo = dati[font_inizio:font_fine]
    dichiarati = set()
    arial = []
    font_nomi = {}
    font_segni = {}
    if sorgente is not None:
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(BytesIO(sorgente))
        for lvl in doc.part.numbering_part.element.iter(qn('w:lvl')):
            text = lvl.find(qn('w:lvlText'))
            fonts = lvl.find(qn('w:rPr'))
            fonts = fonts.find(qn('w:rFonts')) if fonts is not None else None
            if text is not None and fonts is not None:
                glyph, family = text.get(qn('w:val'), ''), fonts.get(qn('w:ascii'), '')
                if len(glyph) == 1 and family:
                    font_segni.setdefault(ord(glyph), set()).add(family)

    for font in re.finditer(rb"\{\\f([0-9]+)\b", catalogo):
        numero = int(font.group(1))
        dichiarati.add(numero)
        voce = catalogo[font.start():_fine_gruppo_rtf(catalogo, font.start())]
        for nome in set().union(*font_segni.values()) if font_segni else ():
            if (b' ' + nome.encode('ascii', errors='ignore') + b';') in voce:
                font_nomi.setdefault(nome, numero)
        if re.search(rb"\bArial;", voce):
            arial.append(numero)
    liste = dati[liste_inizio:liste_fine]
    sostituzioni = []
    for livello in re.finditer(rb"\{\\listlevel\b", liste):
        fine = _fine_gruppo_rtf(liste, livello.start())
        blocco = liste[livello.start():fine]
        for font in re.finditer(rb"\\f([0-9]+)\b", blocco):
            if int(font.group(1)) in dichiarati:
                continue
            target = None
            if re.search(rb"\\levelnfc23\b", blocco):
                glyph = re.search(rb"\\u(-?[0-9]+)\s", blocco)
                families = font_segni.get(int(glyph.group(1)), set()) if glyph else set()
                if len(families) == 1:
                    target = font_nomi.get(next(iter(families)))
                if target is None and arial and re.search(rb"\\u(?:8226|9675|9642)\s", blocco):
                    target = arial[0]
            if target is None:
                raise ValueError("Carattere dell'elenco RTF non dichiarato.")
            sostituzioni.append((livello.start() + font.start(), livello.start() + font.end(),
                                 b"\\f" + str(target).encode("ascii")))
    for inizio, fine, sostituto in reversed(sostituzioni):
        liste = liste[:inizio] + sostituto + liste[fine:]
    return dati[:liste_inizio] + liste + dati[liste_fine:]


def anteprima_stampa_pdf(dati: bytes) -> str:
    """Stampa le pagine renderizzate dal lettore condiviso, senza plugin PDF."""
    import base64
    from pct.rendering_pdf import dimensioni_pagine, pagina_png

    misure = dimensioni_pagine(dati)
    stili = []
    pagine = []
    for misura in misure:
        numero = misura.numero
        larghezza = misura.larghezza / 72 * 25.4
        altezza = misura.altezza / 72 * 25.4
        immagine = pagina_png(dati, numero_pagina=numero, scala=3, predefinita=3)
        codificata = base64.b64encode(immagine).decode('ascii')
        stili.append(f'@page pagina{numero}{{size:{larghezza:.3f}mm {altezza:.3f}mm;margin:0}} .pagina{numero}{{page:pagina{numero};width:{larghezza:.3f}mm;height:{altezza:.3f}mm}} @media print{{.pagina{numero}{{width:{larghezza:.3f}mm!important;height:{altezza:.3f}mm!important;max-width:none!important}}}}')
        pagine.append(f'<section class="pagina pagina{numero}"><img alt="Pagina {numero}" src="data:image/png;base64,{codificata}"></section>')
    return '<!doctype html><html lang="it"><meta charset="utf-8"><title>Anteprima di stampa</title><style>body{margin:0;background:#e5e7eb}.pagina{margin:12px auto;background:#fff;break-after:page}.pagina:last-child{break-after:auto}img{display:block;width:100%;height:100%}@media print{body{background:white}.pagina{margin:0}}</style><style>' + ''.join(stili) + '</style><style>@media screen{.pagina{max-width:100%;height:auto}img{height:auto}}</style><body>' + ''.join(pagine) + '</body></html>'


def esporta_documento_editor(html: str, *, formato: str, titolo: str = "Documento", studio_timbro=None) -> bytes:
    if formato not in {"docx", "rtf", "pdf"}:
        raise ValueError("Formato del documento non consentito.")
    sorgente = html_to_docx(html, titolo=titolo, studio_timbro=studio_timbro)
    if formato == "docx":
        return sorgente
    programma = shutil.which("libreoffice") or shutil.which("soffice")
    if not programma:
        raise RuntimeError("Convertitore locale dei documenti non disponibile.")
    with tempfile.TemporaryDirectory(prefix="iusentra-editor-export-") as cartella:
        root = Path(cartella)
        fonte = root / "documento.docx"
        fonte.write_bytes(sorgente)
        output = root / "output"
        output.mkdir()
        profilo = root / "profilo"
        subprocess.run(
            [programma, "-env:UserInstallation=" + profilo.as_uri(), "--headless", "--convert-to",
             "pdf:writer_pdf_Export" if formato == "pdf" else "rtf", "--outdir", str(output), str(fonte)],
            check=True, timeout=45, capture_output=True,
        )
        prodotto = output / ("documento." + formato)
        if not prodotto.is_file() or prodotto.stat().st_size == 0:
            raise RuntimeError("Il convertitore locale non ha prodotto il documento richiesto.")
        dati = prodotto.read_bytes()
        return _verifica_font_elenchi_rtf(_regioni_rtf(dati, sorgente), sorgente) if formato == "rtf" else dati
